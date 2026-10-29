"""Build KOL总表.xlsx from scratch (layout, styles, data), then let real Excel add the formulas.

Run once. It refuses to overwrite an existing workbook unless --force is given; day-to-day changes go through kol.py.

    python build.py --out KOL总表.xlsx              # an empty sheet
    python build.py --seed my_seed --out X.xlsx     # filled from my_seed.load() (a module next to this file)
    python build.py --out X.xlsx --force
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import kol

FONT = "Microsoft YaHei"
INK, MUTED, FAINT, LINE = "1F2937", "6B7280", "9CA3AF", "E5E7EB"
HEADER_BG, NOTE_BG, SYSTEM_BG = "1E293B", "B45309", "64748B"
CARD_BG, CARD_LINE = "F8FAFC", "E2E8F0"
STAGE_STYLE = {  # fill, font
    "未联系": ("F3F4F6", "6B7280"),
    "已发信": ("E0ECFF", "1D4ED8"),
    "洽谈中": ("FEF3C7", "92400E"),
    "已谈成": ("DCFCE7", "166534"),
    "已婉拒": ("FEE2E2", "991B1B"),
    "退信": ("E5E7EB", "4B5563"),
}
INTENT_COLOR = {"纯CPS": "166534", "CPS+固薪": "6D28D9", "纯固薪": "C2410C", "待定": "9CA3AF"}
MAX_ROW = 5000  # conditional formats reach this far so new rows are covered

MAIN_WIDTHS = {"ID": 8, "KOL": 22, "项目": 9, "语言": 6, "阶段": 9, "意向": 10, "待办方": 8, "下一步": 40,
               "对方报价": 30, "方案/条件": 34, "频道数据": 30, "进展摘要": 44,
               "备注": 24, "平台": 9, "频道": 8, "邮箱": 30, "首次发信": 11, "更新时间": 16, "更新来源": 16,
               "频道链接": 40, "聊天账号": 32}
CENTER_COLS = {"ID", "项目", "语言", "阶段", "意向", "待办方", "平台", "频道", "首次发信"}
LOG_WIDTHS = {"时间": 16, "ID": 8, "KOL": 22, "方向": 6, "类型": 9, "摘要": 76, "渠道": 7, "邮箱": 28, "来源": 18,
              "键": 30}
BATCH_WIDTHS = {"批次": 10, "项目": 9, "语言": 6, "日期": 11, "人数": 6, "主题": 72, "Message-ID": 58}

THIN = Side(style="thin", color=LINE)


def font(size=10, bold=False, color=INK, underline=None, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def setup_table_sheet(ws, headers, widths, rows, table_name, header_colors=None, center=(), row_height=21):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths.get(h, 12)
        c = ws.cell(1, i)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill((header_colors or {}).get(h, HEADER_BG))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    last = max(2, len(rows) + 1)
    for r in range(2, last + 1):
        ws.row_dimensions[r].height = row_height
        for i, h in enumerate(headers, 1):
            c = ws.cell(r, i)
            c.font = font()
            c.alignment = Alignment(horizontal="center" if h in center else "left", vertical="center")
            c.border = Border(bottom=THIN)
    ref = f"A1:{get_column_letter(len(headers))}{last}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 100
    return last


STAGE_ORDER = {"已谈成": 0, "洽谈中": 1, "已发信": 2, "未联系": 3, "已婉拒": 4, "退信": 5}


def build_main(ws, kols):
    headers = kol.MAIN_HEADERS
    rows = []
    kols = sorted(kols, key=lambda k: (STAGE_ORDER[k["stage"]], k["owner"] != "我方", k["id"]))
    for k in kols:
        values = {"ID": k["id"], "KOL": k["kol"], "项目": k["project"], "语言": k["lang"], "阶段": k["stage"],
                  "意向": k["intent"], "待办方": k["owner"], "下一步": k["next"],
                  "对方报价": k["quote"], "方案/条件": k["terms"], "频道数据": k["stats"],
                  "进展摘要": k["summary"], "备注": k["note"], "平台": k["platform"], "频道": None,
                  "邮箱": k["email"], "首次发信": k["first"], "更新时间": k["updated"], "更新来源": k["source"],
                  "频道链接": k["channel"], "聊天账号": k.get("chat", "")}
        rows.append([values.get(h) or None for h in headers])
    colors = {"备注": NOTE_BG, "更新时间": SYSTEM_BG, "更新来源": SYSTEM_BG, "频道链接": SYSTEM_BG}
    last = setup_table_sheet(ws, headers, MAIN_WIDTHS, rows, kol.MAIN_TABLE, colors, CENTER_COLS)
    col = {h: get_column_letter(i) for i, h in enumerate(headers, 1)}
    for r in range(2, last + 1):
        ws[f"{col['KOL']}{r}"].font = font(10, True)
        ws[f"{col['首次发信']}{r}"].number_format = "yyyy-mm-dd"
        ws[f"{col['更新时间']}{r}"].number_format = "yyyy-mm-dd hh:mm"
        for h in ("频道链接", "更新时间", "更新来源", "ID"):
            ws[f"{col[h]}{r}"].font = font(9, False, FAINT)
        ws[f"{col['频道']}{r}"].font = font(10, False, "2563EB", underline="single")
        ws[f"{col['邮箱']}{r}"].font = font(10, False, MUTED)
        ws[f"{col['聊天账号']}{r}"].font = font(10, False, MUTED)
    ws.freeze_panes = "C2"

    # stage pills
    rng = f"{col['阶段']}2:{col['阶段']}{MAX_ROW}"
    for stage, (bg, fg) in STAGE_STYLE.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{stage}"'], stopIfTrue=True,
                                                      fill=fill(bg), font=Font(color=fg, bold=True)))
    rng = f"{col['意向']}2:{col['意向']}{MAX_ROW}"
    for intent, fg in INTENT_COLOR.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{intent}"'], stopIfTrue=True,
                                                      font=Font(color=fg, bold=intent != "待定")))
    rng = f"{col['待办方']}2:{col['待办方']}{MAX_ROW}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"我方"'], stopIfTrue=True,
                                                  fill=fill("FEE2E2"), font=Font(color="B91C1C", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"对方"'], stopIfTrue=True,
                                                  font=Font(color=FAINT)))
    e = col["阶段"]
    # dead rows fade out
    ws.conditional_formatting.add(f"A2:{col['频道链接']}{MAX_ROW}", FormulaRule(
        formula=[f'OR(${e}2="退信",${e}2="已婉拒")'], font=Font(color="9CA3AF")))
    ws.sheet_properties.tabColor = "1D4ED8"


def build_log(ws, logs):
    rows = [[l["time"], l["id"], l["kol"], l["dir"], l["type"], l["summary"], l["channel"], l["email"], l["source"],
             l["key"]] for l in logs]
    last = setup_table_sheet(ws, kol.LOG_HEADERS, LOG_WIDTHS, rows, kol.LOG_TABLE,
                             {"来源": SYSTEM_BG, "键": SYSTEM_BG}, {"ID", "方向", "类型", "渠道"})
    for r in range(2, last + 1):
        ws[f"A{r}"].number_format = "yyyy-mm-dd hh:mm"
        ws[f"C{r}"].font = font(10, True)
        for c in ("B", "I", "J"):
            ws[f"{c}{r}"].font = font(9, False, FAINT)
    # entries known only by date show without a fake 00:00
    dxf = DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode="yyyy-mm-dd"))
    ws.conditional_formatting.add(f"A2:A{MAX_ROW}", Rule(type="expression", dxf=dxf,
                                                         formula=["AND(ISNUMBER($A2),MOD($A2,1)=0)"]))
    ws.conditional_formatting.add(f"D2:D{MAX_ROW}", CellIsRule(operator="equal", formula=['"收"'],
                                                                font=Font(color="1D4ED8", bold=True)))
    ws.conditional_formatting.add(f"D2:D{MAX_ROW}", CellIsRule(operator="equal", formula=['"发"'],
                                                                font=Font(color="15803D", bold=True)))
    ws.conditional_formatting.add(f"A2:H{MAX_ROW}", FormulaRule(
        formula=['OR($E2="首封邀约",$E2="自动回复",$E2="退信")'], font=Font(color="9CA3AF")))
    ws.sheet_properties.tabColor = "0F766E"


def build_batches(ws, batches):
    rows = [[b["batch"], b["project"], b["lang"], b["date"], b["count"], b["subject"], b["key"]] for b in batches]
    last = setup_table_sheet(ws, kol.BATCH_HEADERS, BATCH_WIDTHS, rows, kol.BATCH_TABLE, {"Message-ID": SYSTEM_BG},
                             {"项目", "语言", "日期", "人数"})
    for r in range(2, last + 1):
        ws[f"D{r}"].number_format = "yyyy-mm-dd"
        ws[f"G{r}"].font = font(9, False, FAINT)
    ws.sheet_properties.tabColor = "64748B"


def build_guide(ws):
    ws.sheet_view.showGridLines = False
    for letter, w in {"A": 2, "B": 14, "C": 70, "D": 4, "E": 4, "F": 4, "G": 4,
                      "H": 12, "I": 12, "J": 12, "K": 14, "L": 14}.items():
        ws.column_dimensions[letter].width = w
    ws["B2"] = "使用说明"
    ws["B2"].font = font(18, True)
    lines = [
        "kol-mail 每收到一封 KOL 回信、每发出一封回复，都会马上在这里记一笔，并把待办方切到对应的一方。",
        "工作日 9:40 和 17:00 的回信汇总任务读完新邮件后，会更新阶段、意向、报价、方案和下一步。",
        "在 Claude 会话里处理 KOL 的事时，AI 会按 kol-sheet skill 用同一个工具改表。",
        "表开着也能写。改动会直接出现在窗口里并自动保存；你正在单元格里打字时，AI 会等你打完再写。",
        "备注列留给你自己写，AI 不会改，你要求时才会在后面追加。",
        "每天第一次写入前会自动备份到 kol-sheet\\backups 文件夹，保留最近 14 天。",
        "请不要改各页表头的列名，工具靠列名找位置。加新列、调列宽、排序、筛选都没问题。",
    ]
    ws["B4"] = "这张表怎么更新"
    ws["B4"].font = font(12, True)
    for i, text in enumerate(lines):
        ws.cell(5 + i, 2, f"{i + 1}.").font = font(10, False, MUTED)
        ws.cell(5 + i, 3, text).font = font(10)

    def block(top, title, rows):
        ws.cell(top, 2, title).font = font(12, True)
        for j, (k, v) in enumerate(rows, 1):
            a, b = ws.cell(top + j, 2, k), ws.cell(top + j, 3, v)
            a.font, b.font = font(10, True), font(10)
            a.border = b.border = Border(bottom=THIN)
            if k in STAGE_STYLE:
                a.fill, a.font = fill(STAGE_STYLE[k][0]), font(10, True, STAGE_STYLE[k][1])
            if k in INTENT_COLOR:
                a.font = font(10, True, INTENT_COLOR[k])
        return top + len(rows) + 2

    nxt = block(13, "阶段", [
        ("未联系", "还没发过信，比如找不到邮箱或决定跳过"),
        ("已发信", "首封邀约已发出，对方还没回"),
        ("洽谈中", "对方回过信，正在谈形式和价格"),
        ("已谈成", "双方就付费模式和金额（纯CPS 则是佣金方案）达成一致"),
        ("已婉拒", "对方明确不做，或我方决定放弃"),
        ("退信", "邮件没送到"),
    ])
    nxt = block(nxt, "意向", [
        ("纯固薪", "只要固定费用，拒绝佣金或完全没提佣金"),
        ("CPS+固薪", "要固定费用加佣金，或只有在有保底时才接受 CPS"),
        ("纯CPS", "同意只按佣金或分成合作"),
        ("待定", "婉拒、只在提问、只发了资料没表态，或立场不清楚"),
    ])

    # dropdown sources
    ws["H4"], ws["I4"], ws["J4"] = "阶段列表", "意向列表", "待办方列表"
    for i, v in enumerate(kol.STAGES):
        ws.cell(5 + i, 8, v)
    for i, v in enumerate(kol.INTENTS):
        ws.cell(5 + i, 9, v)
    for i, v in enumerate(kol.OWNERS):
        ws.cell(5 + i, 10, v)
    ws["K4"], ws["K5"] = "项目", kol.DEFAULT_PROJECT
    t = Table(displayName=kol.PROJECT_TABLE, ref="K4:K5")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws["L4"] = "看板项目选项"
    for c in ("H4", "I4", "J4", "K4", "L4"):
        ws[c].font = font(9, True, MUTED)
    for r in range(5, 20):
        for c in "HIJKL":
            ws[f"{c}{r}"].font = font(9, False, MUTED)
    ws["H2"] = "下拉框的选项来源，加新项目时 kol.py 会自动补进“项目”列"
    ws["H2"].font = font(9, False, FAINT, italic=True)
    ws.sheet_properties.tabColor = "B45309"


def card(ws, cols, label, caption_row=7):
    first, last = cols[0], cols[-1]
    for r in (5, 6, 7):
        if first != last:
            ws.merge_cells(f"{first}{r}:{last}{r}")
    for r in (5, 6, 7):
        for c in cols:
            cell = ws[f"{c}{r}"]
            cell.fill = fill(CARD_BG)
            cell.border = Border(left=Side(style="thin", color=CARD_LINE) if c == first else None,
                                 right=Side(style="thin", color=CARD_LINE) if c == last else None,
                                 top=Side(style="thin", color=CARD_LINE) if r == 5 else None,
                                 bottom=Side(style="thin", color=CARD_LINE) if r == 7 else None)
    ws[f"{first}5"] = label
    ws[f"{first}5"].font = font(9, False, MUTED)
    ws[f"{first}5"].alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    ws[f"{first}6"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"{first}7"].alignment = Alignment(horizontal="left", vertical="top", indent=1)
    ws[f"{first}7"].font = font(9, False, FAINT)


def build_dashboard(ws):
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 22, "C": 11, "D": 11, "E": 11, "F": 11, "G": 3, "H": 22, "I": 11, "J": 11, "K": 11,
              "L": 11, "M": 3}
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w
    heights = {1: 10, 2: 34, 3: 20, 4: 12, 5: 22, 6: 40, 7: 20, 8: 14, 9: 26, 10: 22}
    for r, h in heights.items():
        ws.row_dimensions[r].height = h
    for r in range(11, 55):
        ws.row_dimensions[r].height = 21

    ws["B2"] = "KOL 建联总表"
    ws["B2"].font = font(20, True)
    ws["B2"].alignment = Alignment(vertical="center")
    ws["B3"].font = font(9, False, MUTED)
    ws["K2"] = "项目"
    ws["K2"].font = font(9, False, MUTED)
    ws["K2"].alignment = Alignment(horizontal="right", vertical="center")
    ws["L2"] = "全部"
    ws["L2"].font = font(11, True, "1D4ED8")
    ws["L2"].fill = fill("EFF6FF")
    ws["L2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["L2"].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    cards = [(["B"], "总 KOL", INK), (["C", "D"], "已触达", INK), (["E", "F"], "回复率", "1D4ED8"),
             (["H"], "洽谈中", "92400E"), (["I", "J"], "已谈成", "166534"), (["K", "L"], "待我方处理", "B91C1C")]
    for cols, label, color in cards:
        card(ws, cols, label)
        ws[f"{cols[0]}6"].font = font(24, True, color)
    ws["E6"].number_format = "0%"

    def section(cell, title):
        ws[cell] = title
        ws[cell].font = font(12, True)
        ws[cell].alignment = Alignment(vertical="bottom")

    def header_row(row, cells):
        for c, text in cells.items():
            ws[f"{c}{row}"] = text
        for c in cells:
            cell = ws[f"{c}{row}"]
            cell.font = font(9, True, MUTED)
            cell.border = Border(bottom=Side(style="thin", color="CBD5E1"))
            cell.alignment = Alignment(horizontal="left", vertical="center")

    def body(rng_rows, cols, fmt=None):
        for r in rng_rows:
            for c in cols:
                cell = ws[f"{c}{r}"]
                cell.font = font(10)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = Border(bottom=THIN)
                if fmt and c in fmt:
                    cell.number_format = fmt[c]

    # stage distribution
    section("B9", "阶段分布")
    header_row(10, {"B": "阶段", "C": "人数", "D": "占比", "E": "", "F": ""})
    for i, stage in enumerate(kol.STAGES):
        r = 11 + i
        ws[f"B{r}"] = stage
        ws.merge_cells(f"D{r}:F{r}")
    body(range(11, 17), "BCDEF", {"D": "0%"})
    for i, stage in enumerate(kol.STAGES):
        ws[f"B{11 + i}"].font = font(10, True, STAGE_STYLE[stage][1])
    ws.conditional_formatting.add("D11:D16", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                         end_value=1, color="93C5FD", showValue=True))
    # intent distribution
    section("B19", "付费意向（已回复的 KOL）")
    header_row(20, {"B": "意向", "C": "人数", "D": "占比", "E": "", "F": ""})
    for i, intent in enumerate(["纯CPS", "CPS+固薪", "纯固薪", "待定"]):
        r = 21 + i
        ws[f"B{r}"] = intent
        ws.merge_cells(f"D{r}:F{r}")
    body(range(21, 25), "BCDEF", {"D": "0%"})
    for i, intent in enumerate(["纯CPS", "CPS+固薪", "纯固薪", "待定"]):
        ws[f"B{21 + i}"].font = font(10, True, INTENT_COLOR[intent])
    ws.conditional_formatting.add("D21:D24", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                         end_value=1, color="C4B5FD", showValue=True))
    # by language (spills from H11)
    section("H9", "按语言")
    header_row(10, {"H": "语言", "I": "触达", "J": "回复", "K": "回复率", "L": "谈成"})
    body(range(11, 26), "HIJKL", {"K": "0%"})
    ws.conditional_formatting.add("K11:K25", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                         end_value=1, color="93C5FD", showValue=True))
    # to-do lists
    section("B27", "待我方处理")
    header_row(28, {"B": "KOL", "C": "阶段", "D": "意向", "E": "下一步", "F": "", "G": "", "H": "", "I": "",
                    "J": "", "K": "", "L": ""})
    body(range(29, 41), "BCDEFGHIJKL")
    section("B42", "最近沟通")
    header_row(43, {"B": "KOL", "C": "方向", "D": "类型", "E": "时间", "F": "摘要", "G": "", "H": "", "I": "",
                    "J": "", "K": "", "L": ""})
    body(range(44, 54), "BCDEFGHIJKL", {"E": "mm-dd hh:mm"})
    for r in list(range(29, 41)) + list(range(44, 54)):
        ws[f"B{r}"].font = font(10, True)
    for rng in ("C29:C40",):
        for stage, (bg, fg) in STAGE_STYLE.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{stage}"'],
                                                          font=Font(color=fg, bold=True)))
    ws.conditional_formatting.add("E44:E53", Rule(type="expression", formula=["AND(ISNUMBER($E44),MOD($E44,1)=0)"],
        dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=201, formatCode="mm-dd"))))
    ws.conditional_formatting.add("C44:C53", CellIsRule(operator="equal", formula=['"收"'],
                                                        font=Font(color="1D4ED8", bold=True)))
    ws.conditional_formatting.add("C44:C53", CellIsRule(operator="equal", formula=['"发"'],
                                                        font=Font(color="15803D", bold=True)))
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "1E293B"


# ---------------------------------------------------------------- formulas, names and validation, written by Excel

PROJ_MATCH = 'IF(所选项目="全部",TRUE,KOL[项目]=所选项目)'
REACHED = '{"已发信","洽谈中","已谈成","已婉拒"}'
REPLIED = '{"洽谈中","已谈成","已婉拒"}'


def dashboard_formulas() -> dict[str, str]:
    f = {
        "B3": '="数据更新于 "&TEXT(MAX(KOL[更新时间]),"yyyy-mm-dd hh:mm")&"　·　今天 "&TEXT(TODAY(),"yyyy-mm-dd")',
        "B6": "=COUNTIFS(KOL[项目],项目条件)",
        "B7": '="未联系 "&COUNTIFS(KOL[项目],项目条件,KOL[阶段],"未联系")&" · 退信 "&COUNTIFS(KOL[项目],项目条件,KOL[阶段],"退信")',
        "C6": f"=SUM(COUNTIFS(KOL[项目],项目条件,KOL[阶段],{REACHED}))",
        "C7": '="其中首封未回 "&COUNTIFS(KOL[项目],项目条件,KOL[阶段],"已发信")&" 人"',
        "E6": f"=IFERROR(SUM(COUNTIFS(KOL[项目],项目条件,KOL[阶段],{REPLIED}))/C6,0)",
        "E7": f'="已回复 "&SUM(COUNTIFS(KOL[项目],项目条件,KOL[阶段],{REPLIED}))&" 人"',
        "H6": '=COUNTIFS(KOL[项目],项目条件,KOL[阶段],"洽谈中")',
        "H7": '="其中待我方处理 "&COUNTIFS(KOL[项目],项目条件,KOL[阶段],"洽谈中",KOL[待办方],"我方")&" 人"',
        "I6": '=COUNTIFS(KOL[项目],项目条件,KOL[阶段],"已谈成")',
        "I7": '=LET(l,{"纯CPS","CPS+固薪","纯固薪","待定"},n,COUNTIFS(KOL[项目],项目条件,KOL[阶段],"已谈成",KOL[意向],l),t,TEXTJOIN(" · ",TRUE,IF(n>0,l&" "&n,"")),IF(t="","暂无",t))',
        "K6": '=COUNTIFS(KOL[项目],项目条件,KOL[待办方],"我方")',
        "K7": '="需要你回复或推进"',
        "H11": f'=LET(m,{PROJ_MATCH},u,UNIQUE(FILTER(KOL[语言],m*(KOL[语言]<>""),"")),n,COUNTIFS(KOL[语言],u,KOL[项目],项目条件),TAKE(SORTBY(u,n,-1),15))',
        "I11": '=IF(H11#="","",COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件)-COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件,KOL[阶段],"未联系")-COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件,KOL[阶段],"退信"))',
        "J11": '=IF(H11#="","",COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件,KOL[阶段],"洽谈中")+COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件,KOL[阶段],"已谈成")+COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件,KOL[阶段],"已婉拒"))',
        "K11": '=IF(H11#="","",IFERROR(J11#/I11#,0))',
        "L11": '=IF(H11#="","",COUNTIFS(KOL[语言],H11#,KOL[项目],项目条件,KOL[阶段],"已谈成"))',
        "B27": '="待我方处理　"&COUNTIFS(KOL[项目],项目条件,KOL[待办方],"我方")&" 位"&IF(COUNTIFS(KOL[项目],项目条件,KOL[待办方],"我方")>12,"，这里只显示前 12 位","")',
        "B29": f'=LET(m,{PROJ_MATCH},o,IFERROR(MATCH(KOL[阶段],{{"已谈成","洽谈中","已发信","未联系","已婉拒","退信"}},0),9),f,FILTER(HSTACK(KOL[KOL]&"",KOL[阶段]&"",KOL[意向]&"",KOL[下一步]&"",o),m*(KOL[待办方]="我方"),""),IF(ROWS(f)*COLUMNS(f)=1,"（暂无）",TAKE(CHOOSECOLS(SORTBY(f,INDEX(f,0,5),1,INDEX(f,0,1),1),1,2,3,4),12)))',
        "B44": '=LET(k,(LOG[类型]<>"自动回复")*(LOG[类型]<>"首封邀约"),f,FILTER(HSTACK(LOG[KOL]&"",LOG[方向]&"",LOG[类型]&"",LOG[时间],LOG[摘要]&""),k,""),IF(ROWS(f)*COLUMNS(f)=1,"（暂无）",TAKE(SORTBY(f,INDEX(f,0,4),-1),10)))',
    }
    for i in range(6):
        r = 11 + i
        f[f"C{r}"] = f"=COUNTIFS(KOL[项目],项目条件,KOL[阶段],B{r})"
        f[f"D{r}"] = f"=IFERROR(C{r}/$B$6,0)"
    for i in range(4):
        r = 21 + i
        f[f"C{r}"] = f"=SUM(COUNTIFS(KOL[项目],项目条件,KOL[意向],B{r},KOL[阶段],{REPLIED}))"
        f[f"D{r}"] = f"=IFERROR(C{r}/SUM($C$21:$C$24),0)"
    return f


def finalize_with_excel(path: Path):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("Excel.Application")
    pid = kol.process_of(app)
    app.Visible = False
    app.DisplayAlerts = False
    try:
        wb = app.Workbooks.Open(str(path))
        names = {
            "所选项目": "=看板!$L$2",
            "项目条件": '=IF(看板!$L$2="全部","<>"&CHAR(1),看板!$L$2)',  # "*" would skip blank cells
            "阶段列表": f"=说明!$H$5:$H${4 + len(kol.STAGES)}",
            "意向列表": f"=说明!$I$5:$I${4 + len(kol.INTENTS)}",
            "待办方列表": f"=说明!$J$5:$J${4 + len(kol.OWNERS)}",
            "项目列表": "=PROJ[项目]",
            "看板项目选项": "=说明!$L$5#",
        }
        guide = wb.Worksheets("说明")
        guide.Range("L5").Formula2 = '=VSTACK("全部",PROJ[项目])'
        for name, ref in names.items():
            wb.Names.Add(name, ref)

        main = wb.Worksheets(kol.MAIN_SHEET)
        lo = main.ListObjects(kol.MAIN_TABLE)
        lo.ListColumns("频道").DataBodyRange.Formula = '=IF([@频道链接]="","",HYPERLINK([@频道链接],"打开 ↗"))'
        for column, source in (("阶段", "=阶段列表"), ("意向", "=意向列表"), ("待办方", "=待办方列表"), ("项目", "=项目列表")):
            rng = lo.ListColumns(column).DataBodyRange
            rng.Validation.Delete()
            rng.Validation.Add(3, 1, 1, source)  # xlValidateList, xlValidAlertStop, xlBetween
            rng.Validation.IgnoreBlank = True
            rng.Validation.InCellDropdown = True
        log = wb.Worksheets(kol.LOG_SHEET).ListObjects(kol.LOG_TABLE)
        rng = log.ListColumns("方向").DataBodyRange
        rng.Validation.Delete()
        rng.Validation.Add(3, 1, 1, "收,发")

        dash = wb.Worksheets("看板")
        dash.Activate()  # Excel can crash writing dynamic-array formulas into a sheet that is not active
        for cell, formula in dashboard_formulas().items():
            dash.Range(cell).Formula2 = formula
        v = dash.Range("L2").Validation
        v.Delete()
        v.Add(3, 1, 1, "=看板项目选项")

        app.CalculateFull()
        problems = []
        for ws in wb.Worksheets:
            try:
                errs = ws.UsedRange.SpecialCells(-4123, 16)  # formulas that evaluate to an error
                problems.append(f"{ws.Name}!{errs.Address}")
            except Exception:
                pass
        for ws in wb.Worksheets:
            ws.Activate()
            ws.Range("A1").Select()
            if ws.Name == kol.LOG_SHEET:  # the log grows downwards: open it at the latest entries
                last = log.Range.Rows.Count
                app.Goto(ws.Range(f"A{max(2, last - 20)}"), True)
                ws.Range(f"A{last}").Select()
        dash.Activate()
        wb.Save()
        wb.Close(False)
        return problems
    finally:
        app.Quit()
        del app
        kol.make_sure_exited(pid)
        pythoncom.CoUninitialize()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(kol.WORKBOOK))
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--seed", default="", help="module next to this file with a load() -> (kols, logs, batches)")
    a = ap.parse_args()
    out = Path(a.out).resolve()
    if out.exists() and not a.force:
        sys.exit(f"{out} 已存在。日常改动请用 kol.py；确实要重建请加 --force（会覆盖现有数据）")
    kols, logs, batches = [], [], []  # an empty sheet; a one-off migration module can still seed it
    if a.seed:
        seed = __import__(a.seed)
        kols, logs, batches = seed.load()

    wb = Workbook()
    dash = wb.active
    dash.title = "看板"
    main_ws = wb.create_sheet(kol.MAIN_SHEET)
    log_ws = wb.create_sheet(kol.LOG_SHEET)
    batch_ws = wb.create_sheet(kol.BATCH_SHEET)
    guide_ws = wb.create_sheet(kol.DICT_SHEET)
    build_dashboard(dash)
    build_main(main_ws, kols)
    build_log(log_ws, logs)
    build_batches(batch_ws, batches)
    build_guide(guide_ws)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    problems = finalize_with_excel(out)
    print(f"已生成 {out}：{len(kols)} 位 KOL，{len(logs)} 条沟通记录，{len(batches)} 个发信批次")
    if problems:
        print("以下单元格公式出错：", "; ".join(problems))
        sys.exit(1)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
