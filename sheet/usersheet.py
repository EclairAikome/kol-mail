"""KOL 总表同步，给用自己表格的同事用（主人那张固定格式的总表还是走 kol.py）。

每个人的总表格式不一样，这里不改结构、不加列、不加工作表、不动格式，只往“列对应”里对上的几列写值：
邮箱（必须）、KOL 名字、频道链接、最近联系时间、待办方、进展摘要。没对上的列就不写那一项。

- 每收到或发出一封 KOL 邮件：按邮箱找到那一行，写最近联系时间、轮到谁回、一句话进展。比表里已有的更旧的邮件
  不会盖掉新情况。表里没有这个邮箱时在末尾加一行，格式照抄上一行（在 Excel 表格对象里就让表格自己扩一行）。
- 页面打开一封信时显示的频道快捷链接，来自“邮箱”和“频道链接”两列（list 命令，直接读文件，不启动 Excel）。

Windows 用 Excel COM：表在 Excel 里开着就写进打开的窗口，没开就在后台开一个看不见的 Excel 写完保存再关。
Mac 用 AppleScript 控制 Excel（macexcel.py），开着就写进窗口，没开就让 Excel 在后台打开、写完、保存、关上。Excel 正忙或文件被占用时改动排进 pending.jsonl，
之后任何一次调用都会先补写。命令和退出码与 kol.py 一样，app.py 不用区分：
0 写好了，2 用法或列对应有问题，3 排队了，4 收下了但有几条没写成。

列对应由页面的设置传进来：环境变量 KOL_SHEET_MAP，JSON {"sheet": 工作表名, "header_row": 表头行号,
"columns": {"email": 表头文字, ...}}。按表头文字找列，同事挪动列的顺序也不受影响。
"""

from __future__ import annotations

import argparse
import contextlib
import datetime as dt
import gc
import json
import os
import re
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))  # kol.py (the Excel COM helpers) sits next to this file
WORKBOOK = Path(os.environ.get("KOL_SHEET_PATH") or "")
STATE_DIR = Path(os.environ.get("KOL_SHEET_STATE") or HERE)
PENDING = STATE_DIR / "pending.jsonl"
FAILED = STATE_DIR / "pending.failed.jsonl"
STATE_FILE = STATE_DIR / "usersheet_state.json"
BACKUPS = STATE_DIR / "backups"
KEEP_BACKUPS = 14
KEEP_KEYS = 5000
BUSY_WAIT_SECONDS = 60
UNAVAILABLE_WAIT_SECONDS = 8
STALE_TOLERANCE = dt.timedelta(minutes=10)  # covers senders whose clocks run slow
DEFAULT_DATE_FORMAT = "yyyy-mm-dd hh:mm"

# field -> what the settings page calls it
FIELDS = {
    "email": "邮箱",
    "name": "KOL 名字",
    "channel": "频道链接",
    "last_contact": "最近联系时间",
    "owner": "待办方",
    "summary": "进展摘要",
}
REQUIRED = ("email",)
# header words that usually mean the field (compared lower-case, spaces and punctuation removed)
SYNONYMS = {
    "email": ["邮箱", "邮件", "邮箱地址", "联系邮箱", "电子邮件", "email", "e-mail", "mail", "emailaddress", "contactemail"],
    "name": ["kol", "kol名", "kol名字", "kol名称", "名字", "名称", "姓名", "达人", "博主", "主播", "红人", "频道名",
             "频道名称", "账号", "账号名", "昵称", "name", "creator", "influencer", "channelname", "streamer", "youtuber",
             "handle"],
    "channel": ["频道链接", "频道", "链接", "主页", "主页链接", "频道地址", "channel", "channellink", "channelurl", "link",
                "url", "youtube", "twitch", "youtube链接", "twitch链接", "profile", "profilelink"],
    "last_contact": ["最近联系时间", "最近联系", "最后联系", "最后联系时间", "联系时间", "更新时间", "最近更新", "跟进时间",
                     "最近跟进", "lastcontact", "lastcontacted", "lastcontactdate", "lastupdate", "updated", "lastemail",
                     "lastreply", "date"],
    "owner": ["待办方", "待办", "轮到", "轮到谁", "谁回", "等谁回", "ball", "ballin", "waitingon", "whosturn",
              "whosenext"],
    "summary": ["进展摘要", "进展", "最新进展", "跟进情况", "沟通情况", "沟通进展", "最新情况", "summary", "progress",
                "lastmessage", "latestupdate"],
}
OURS = ["我方", "我们", "我", "己方", "自己", "本方", "me", "us", "we", "ours", "oursside", "myturn", "our turn"]
THEIRS = ["对方", "kol", "达人", "博主", "红人", "主播", "他们", "them", "they", "theirs", "creator", "influencer",
          "theirturn", "their turn"]
LINK = re.compile(r"https?://\S+|(?:www\.)?(?:youtube\.com|youtu\.be|twitch\.tv|tiktok\.com|t\.me)/\S+", re.IGNORECASE)


class UsageError(Exception):
    pass


class Unavailable(Exception):
    pass


# ---------------------------------------------------------------- small helpers

def norm_header(text) -> str:
    return re.sub(r"[\s_\-:：/（）()\[\]【】.·]+", "", str(text or "")).lower()


def split_emails(value) -> list[str]:
    """Same split as kol.py and app.py: "<a@x.com>、b@y.com" holds two addresses."""
    return [e.strip("<>()（）") for e in re.split(r"[;,，；、/\s]+", str(value or "").lower()) if "@" in e]


def now() -> dt.datetime:
    return dt.datetime.now().replace(microsecond=0)


def parse_when(value) -> dt.datetime:
    """Local naive time: Excel cells carry no time zone."""
    d = dt.datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if d.tzinfo is not None:
        d = d.astimezone().replace(tzinfo=None)
    return d.replace(microsecond=0)


def from_serial(x) -> dt.datetime | None:
    if isinstance(x, dt.datetime):
        return x.replace(tzinfo=None)
    if isinstance(x, dt.date):
        return dt.datetime(x.year, x.month, x.day)
    if isinstance(x, (int, float)) and 1 < x < 2958466:
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(x))
    if isinstance(x, str) and x.strip():
        with contextlib.suppress(ValueError):
            return parse_when(x.strip().replace("/", "-"))
    return None


def to_serial(d: dt.datetime) -> float:
    return (d - dt.datetime(1899, 12, 30)).total_seconds() / 86400


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def safe_text(value: str) -> str:
    """What to put into a cell so Excel stores it as text: never a formula, never turned into a number or date."""
    value = str(value or "").replace("\r", " ").replace("\n", " ").strip()[:500]
    if value[:1] in ("=", "+", "-", "@") or re.fullmatch(r"[\d\s.,:/%-]+", value or "x"):
        return "'" + value
    return value


# ---------------------------------------------------------------- column mapping

def load_map(raw: str | None = None) -> dict:
    raw = os.environ.get("KOL_SHEET_MAP", "") if raw is None else raw
    try:
        m = json.loads(raw) if raw else {}
    except ValueError:
        raise UsageError("列对应设置读不出来，请在页面的设置里重新对应一次") from None
    if not isinstance(m, dict) or not m.get("sheet") or not isinstance(m.get("columns"), dict):
        raise UsageError("还没有设置列对应，请在页面的设置里打开“KOL 总表”对应一下")
    cols = {k: str(v) for k, v in m["columns"].items() if k in FIELDS and str(v or "").strip()}
    if "email" not in cols:
        raise UsageError("列对应里没有“邮箱”列，同步要靠邮箱找到对应的那一行")
    try:
        header_row = int(m.get("header_row") or 1)
    except (TypeError, ValueError):
        header_row = 1
    return {"sheet": str(m["sheet"]), "header_row": max(1, header_row), "columns": cols}


def find_header_row(rows: list[list], limit: int = 15) -> int:
    """1-based row that looks most like a header: many short text cells, ideally with words we know."""
    best, best_score = 1, -1.0
    known = {norm_header(w) for words in SYNONYMS.values() for w in words}
    for i, row in enumerate(rows[:limit]):
        cells = [text(v) for v in row if text(v)]
        if not cells:
            continue
        texts = [c for c in cells if not from_serial(c) and len(c) <= 40 and "@" not in c and "://" not in c]
        hits = sum(1 for c in texts if norm_header(c) in known)
        score = len(texts) + 3 * hits - (len(cells) - len(texts))
        if score > best_score:
            best, best_score = i + 1, score
    return best


def column_profile(rows: list[list], header_row: int, col: int, sample: int = 300) -> dict:
    body = [text(r[col - 1]) if col - 1 < len(r) else "" for r in rows[header_row:header_row + sample]]
    filled = [v for v in body if v]
    return {"filled": len(filled),
            "emails": sum(1 for v in filled if split_emails(v)),
            "links": sum(1 for v in filled if LINK.search(v)),
            "dates": sum(1 for v in filled if from_serial(v)),
            "samples": filled[:3]}


def guess_map(rows: list[list], header_row: int) -> dict:
    """field -> header text, from header words first and the column's content second; one column per field."""
    headers = rows[header_row - 1] if header_row - 1 < len(rows) else []
    cols = {c + 1: text(h) for c, h in enumerate(headers) if text(h)}
    profiles = {c: column_profile(rows, header_row, c) for c in cols}
    scores: list[tuple[float, str, int]] = []
    for field, words in SYNONYMS.items():
        normed = [norm_header(w) for w in words]
        for c, h in cols.items():
            key, p = norm_header(h), profiles[c]
            s = 0.0
            if key in normed:
                s = 10 - normed.index(key) * 0.05
            elif any(w and len(w) >= 2 and w in key for w in normed):
                s = 5
            if field == "email":
                s += 6 * p["emails"] / max(p["filled"], 1)
            elif field == "channel":
                s += 4 * p["links"] / max(p["filled"], 1) if s else 0
            elif field == "last_contact":
                s += 2 * p["dates"] / max(p["filled"], 1) if s else 0
            if p["emails"] and field != "email":
                s -= 4
            if s >= 5:
                scores.append((s, field, c))
    out: dict[str, str] = {}
    used: set[int] = set()
    for s, field, c in sorted(scores, reverse=True):
        if field not in out and c not in used:
            out[field], _ = cols[c], used.add(c)
    return out


# ---------------------------------------------------------------- workbooks

class Grid:
    """What the sync needs from a workbook. Rows and columns are 1-based sheet coordinates."""
    mode = ""

    def sheet_names(self) -> list[str]:
        raise NotImplementedError

    def rows(self, sheet: str) -> list[list]:
        """Every value from A1 to the last used cell; rows[r - 1][c - 1] is cell (r, c)."""
        raise NotImplementedError

    def set_text(self, sheet: str, r: int, c: int, value: str):
        raise NotImplementedError

    def set_date(self, sheet: str, r: int, c: int, value: dt.datetime, header_row: int):
        raise NotImplementedError

    def append_row(self, sheet: str, header_row: int, last_row: int, free_row: int) -> int:
        """An empty row for a new KOL, formatted like `last_row` (the last data row); `free_row` is the first row
        below it that is empty in every column. Returns the row number used."""
        raise NotImplementedError

    def save(self):
        pass

    def backup_once_a_day(self):
        pass


class XlsxGrid(Grid):
    """openpyxl on the file itself. Used to read (list, inspect, new sheets) and in tests; the sync never writes a
    colleague's file this way, because openpyxl drops what it does not understand (charts, images, some formats)."""
    mode = "file"

    def __init__(self, path: Path, write: bool = False):
        from openpyxl import load_workbook
        self.path = path
        try:
            self.wb = load_workbook(path, data_only=not write, read_only=False)
        except FileNotFoundError:
            raise UsageError(f"找不到总表：{path}") from None
        except PermissionError:
            raise Unavailable("总表文件现在打不开（被别的程序占用）") from None
        except Exception as e:  # noqa: BLE001  not an xlsx, encrypted, broken
            raise UsageError(f"读不了这个文件（要 .xlsx 格式，不能加密）：{type(e).__name__}: {e}") from None

    def sheet_names(self):
        return list(self.wb.sheetnames)

    def _ws(self, sheet):
        if sheet not in self.wb.sheetnames:
            raise UsageError(f"总表里没有叫“{sheet}”的工作表了（被改名或删掉了？），请在设置里重新对应")
        return self.wb[sheet]

    def rows(self, sheet):
        ws = self._ws(sheet)
        return [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True)]

    def set_text(self, sheet, r, c, value):
        cell = self._ws(sheet).cell(r, c)
        cell.value = value[1:] if value.startswith("'") else value
        if value.startswith("'"):
            cell.data_type = "s"

    def set_date(self, sheet, r, c, value, header_row):
        ws = self._ws(sheet)
        cell = ws.cell(r, c)
        if cell.number_format == "@":
            cell.value = f"{value:%Y-%m-%d %H:%M}"
            return
        if cell.number_format in ("General", None):
            cell.number_format = next((ws.cell(k, c).number_format for k in range(r - 1, header_row, -1)
                                       if ws.cell(k, c).number_format not in ("General", None)), DEFAULT_DATE_FORMAT)
        cell.value = value

    def append_row(self, sheet, header_row, last_row, free_row):
        from copy import copy
        ws = self._ws(sheet)
        new = free_row
        if last_row > header_row:
            for c in range(1, ws.max_column + 1):
                src, dst = ws.cell(last_row, c), ws.cell(new, c)
                if src.has_style:
                    dst._style = copy(src._style)
        for t in ws.tables.values():  # grow an Excel table that ends right above
            from openpyxl.utils.cell import range_boundaries
            c1, r1, c2, r2 = range_boundaries(t.ref)
            if r1 == header_row and r2 == last_row and new == last_row + 1:
                from openpyxl.utils import get_column_letter
                t.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{new}"
        return new

    def save(self):
        self.wb.save(self.path)


class ExcelGrid(Grid):
    """Excel through COM (Windows). Open in Excel: written into that window, which the colleague sees at once and
    saves with Ctrl+S as usual (we save too). Not open: a hidden Excel opens it, writes, saves and quits."""

    def __init__(self, path: Path, write: bool = True):
        self.path, self.write = path, write
        self.wb = self.app = None
        self.pid = None
        self.mode = ""

    def __enter__(self):
        import pythoncom
        pythoncom.CoInitialize()
        try:
            self._connect()
            return self
        except BaseException:
            self.__exit__(None, None, None)
            raise

    def _connect(self):
        import win32com.client
        import kol
        if not self.path.exists():
            raise UsageError(f"找不到总表：{self.path}")
        self.wb = kol.find_open_workbook(self.path)
        if self.wb is not None:
            app = self.wb.Application
            self.mode = "live" if (app.Visible or app.UserControl) else "orphan"
            if self.mode == "orphan":
                self.app, self.pid = app, kol.process_of(app)
        else:
            if kol.file_locked(self.path):
                raise Unavailable("总表被别的程序占用了（不在本机 Excel 的已打开列表里，可能是受保护视图或别的软件）")
            self.app = win32com.client.DispatchEx("Excel.Application")
            self.pid = kol.process_of(self.app)
            self.app.Visible = False
            self.app.DisplayAlerts = False
            self.app.ScreenUpdating = False
            self.app.EnableEvents = False
            self.app.AskToUpdateLinks = False
            self.wb = self.app.Workbooks.Open(str(self.path), 0, False)
            self.mode = "headless"
        if self.write and self.wb.ReadOnly:
            raise Unavailable("总表在 Excel 里是以只读方式打开的，存不进去。请关掉后正常打开")

    def __exit__(self, exc_type, exc, tb):
        import pythoncom
        import kol
        ours = self.mode in ("headless", "orphan") or (not self.mode and self.app is not None)
        if ours and self.mode == "orphan":
            with contextlib.suppress(Exception):
                if self.app.Workbooks.Count > 1:
                    self.wb.Close(False)
                    ours = False
        try:
            if ours:
                with contextlib.suppress(Exception):
                    if self.wb is not None:
                        self.wb.Close(False)
                with contextlib.suppress(Exception):
                    self.app.Quit()
        finally:
            self.wb = self.app = None
            gc.collect()
            if ours:
                kol.make_sure_exited(self.pid)
            pythoncom.CoUninitialize()

    def _ws(self, sheet):
        try:
            return self.wb.Worksheets(sheet)
        except Exception:
            raise UsageError(f"总表里没有叫“{sheet}”的工作表了（被改名或删掉了？），请在设置里重新对应") from None

    def sheet_names(self):
        return [self.wb.Worksheets(i).Name for i in range(1, self.wb.Worksheets.Count + 1)]

    def rows(self, sheet):
        ws = self._ws(sheet)
        used = ws.UsedRange
        r0, c0 = used.Row, used.Column
        v = used.Value2
        if not isinstance(v, tuple):
            v = ((v,),)
        width = c0 - 1 + max((len(r) for r in v), default=0)
        out = [[None] * width for _ in range(r0 - 1)]
        for row in v:
            out.append([None] * (c0 - 1) + list(row))
        return out

    def set_text(self, sheet, r, c, value):
        self._ws(sheet).Cells(r, c).Value2 = value

    def set_date(self, sheet, r, c, value, header_row):
        ws = self._ws(sheet)
        cell = ws.Cells(r, c)
        if str(cell.NumberFormat) == "@":  # a text column: write the date as text, as the colleague types it
            cell.Value2 = f"{value:%Y-%m-%d %H:%M}"
            return
        if str(cell.NumberFormat) == "General":
            fmt = DEFAULT_DATE_FORMAT
            for k in range(r - 1, header_row, -1):
                f = str(ws.Cells(k, c).NumberFormat)
                if f != "General" and ws.Cells(k, c).Value2 is not None:
                    fmt = f
                    break
            cell.NumberFormat = fmt
        cell.Value2 = to_serial(value)

    def append_row(self, sheet, header_row, last_row, free_row):
        ws = self._ws(sheet)
        for i in range(1, ws.ListObjects.Count + 1):  # data kept in an Excel table: let the table grow
            lo = ws.ListObjects(i)
            if lo.HeaderRowRange.Row == header_row:
                body = lo.DataBodyRange
                end = body.Row + body.Rows.Count - 1 if body is not None else header_row
                if end > last_row and free_row <= end:
                    return free_row  # an empty row inside the table already
                lo.ListRows.Add()  # inserts inside the table: whatever sits below moves down, nothing is overwritten
                return lo.ListRows(lo.ListRows.Count).Range.Row
        new = free_row
        if last_row > header_row:
            ws.Rows(last_row).Copy(ws.Rows(new))  # formats, validation and formulas, without the clipboard
            with contextlib.suppress(Exception):  # no constants in the row: SpecialCells raises
                ws.Rows(new).SpecialCells(2).ClearContents()  # xlCellTypeConstants: keep formulas, drop values
        return new

    def save(self):
        self.wb.Save()

    def backup_once_a_day(self):
        BACKUPS.mkdir(parents=True, exist_ok=True)
        target = BACKUPS / f"{self.path.stem}_{dt.date.today():%Y%m%d}{self.path.suffix}"
        if not target.exists():
            self.wb.SaveCopyAs(str(target))
        for f in sorted(BACKUPS.glob(f"{self.path.stem}_*{self.path.suffix}"))[:-KEEP_BACKUPS]:
            with contextlib.suppress(OSError):
                f.unlink()


@contextlib.contextmanager
def open_grid(write: bool = True):
    if sys.platform == "win32":
        with ExcelGrid(WORKBOOK, write) as g:
            yield g
        return
    if sys.platform == "darwin":
        from macexcel import MacExcelGrid
        with MacExcelGrid(WORKBOOK, write, Unavailable, UsageError, BACKUPS, KEEP_BACKUPS) as g:
            yield g
        return
    raise Unavailable("这个系统上没有能写 Excel 的办法")


# ---------------------------------------------------------------- the sync itself

OUR_SUMMARY = ("对方来信", "对方回信", "我方回信", "收到回信")  # how every summary this tool writes starts

class UserSheet:
    def __init__(self, grid: Grid, mapping: dict, state: dict):
        self.g, self.m, self.state = grid, mapping, state
        self.sheet, self.hr = mapping["sheet"], mapping["header_row"]
        self.reload()
        self.learn_words()
        missing = [f"{FIELDS[f]}（{h}）" for f, h in mapping["columns"].items() if f not in self.col]
        if missing:
            raise UsageError(f"总表“{self.sheet}”第 {self.hr} 行的表头里找不到这些列：{'、'.join(missing)}。"
                             "列名改过的话，在页面的设置里重新对应一次")

    def reload(self):
        self.rows = self.g.rows(self.sheet)
        headers = self.rows[self.hr - 1] if self.hr - 1 < len(self.rows) else []
        by_text = {norm_header(h): i + 1 for i, h in reversed(list(enumerate(headers))) if text(h)}
        self.col = {f: by_text[norm_header(h)] for f, h in self.m["columns"].items() if norm_header(h) in by_text}

    def cell(self, r: int, field: str):
        c = self.col.get(field)
        if not c or r - 1 >= len(self.rows) or c - 1 >= len(self.rows[r - 1]):
            return None
        return self.rows[r - 1][c - 1]

    def free_row(self, after: int) -> int:
        """First row below `after` with nothing in any column, so a new KOL never lands on a total, a note or a
        pre-numbered row."""
        r = after + 1
        while r - 1 < len(self.rows) and any(text(v) for v in self.rows[r - 1]):
            r += 1
        return r

    def last_row(self) -> int:
        """Last row with anything in a mapped column (not the formatting tail Excel calls "used")."""
        for r in range(len(self.rows), self.hr, -1):
            if any(text(self.cell(r, f)) for f in self.col):
                return r
        return self.hr

    def find(self, email: str) -> int | None:
        for r in range(self.hr + 1, len(self.rows) + 1):
            if email in split_emails(self.cell(r, "email")):
                return r
        return None

    def put(self, r: int, field: str, value):
        c = self.col.get(field)
        if not c:
            return
        if isinstance(value, dt.datetime):
            self.g.set_date(self.sheet, r, c, value, self.hr)
        else:
            self.g.set_text(self.sheet, r, c, safe_text(value))
        while len(self.rows) < r:
            self.rows.append([])
        row = self.rows[r - 1]
        while len(row) < c:
            row.append(None)
        row[c - 1] = value

    def learn_words(self):
        """The words this sheet uses for whose turn it is, read before anything is written (a batch can turn every
        "Them" into "Me") and kept in the state for when the column no longer shows one of them."""
        saved = self.state.get("owner_words") or {}
        found: dict[str, str] = {}
        for r in range(self.hr + 1, len(self.rows) + 1):
            v = text(self.cell(r, "owner"))
            for side, words in (("ours", OURS), ("theirs", THEIRS)):
                if v and side not in found and norm_header(v) in {norm_header(w) for w in words}:
                    found[side] = v
        self.words = {**saved, **found}
        if self.words:
            self.state["owner_words"] = self.words

    def ours_to_change(self, r: int, field: str) -> bool:
        """Only an empty cell or one this sync could have written may be overwritten; a colleague's own text stays."""
        v = self.cell(r, field)
        if not text(v):
            return True
        if field == "last_contact":
            return from_serial(v) is not None
        if field == "owner":
            known = {norm_header(w) for w in OURS + THEIRS} | {norm_header(w) for w in self.words.values()}
            return norm_header(v) in known
        if field == "summary":
            return text(v).lstrip("'").startswith(OUR_SUMMARY)
        return False

    def vocab(self, ours: bool) -> str:
        """The word this sheet already uses for whose turn it is, or 我方/对方."""
        return self.words.get("ours" if ours else "theirs") or ("我方" if ours else "对方")

    # -- operations

    def event(self, op: dict) -> dict:
        email, when, key = op["email"], op["time"], op.get("key") or ""
        if op.get("auto"):  # an automatic reply says nothing about whose turn it is
            return {"status": "skipped", "note": "自动回复不改总表"}
        seen = self.state.setdefault("emails", {}).get(email) or {}
        keys = self.state.setdefault("keys", [])
        repeat = bool(key) and key in keys
        r = self.find(email)
        created = False
        if r is None:
            if repeat:
                return {"status": "skipped", "note": "这封信记过，那一行后来被删了，不再加回去"}
            last = self.last_row()
            r = self.g.append_row(self.sheet, self.hr, last, self.free_row(last))
            self.put(r, "email", email)
            self.put(r, "name", op.get("name") or email.split("@")[0])
            if op.get("channel_url"):
                self.put(r, "channel", op["channel_url"])
            created = True
        if op.get("channel_url") and not created and not text(self.cell(r, "channel")):
            self.put(r, "channel", op["channel_url"])  # a link found in the mail fills an empty cell, never replaces
        latest = from_serial(seen.get("time"))
        in_cell = from_serial(self.cell(r, "last_contact"))
        newest = max([d for d in (latest, in_cell) if d], default=None)
        if repeat:
            if seen.get("key") == key and op.get("summary_given") and self.ours_to_change(r, "summary"):
                self.put(r, "summary", op["summary"])  # the Chinese gist of the newest mail
                return {"status": "ok", "row": r, "updated": ["summary"]}
            return {"status": "skipped", "row": r, "note": "这封信已经记过"}
        if key:
            keys.append(key)
            del keys[:-KEEP_KEYS]
        if newest and when < newest - STALE_TOLERANCE:
            return {"status": "ok", "row": r, "created": created, "note": "比表里的情况旧，没有覆盖"}
        kept = []
        for field, value in (("last_contact", when), ("owner", self.vocab(ours=op["dir"] == "in")),
                             ("summary", op["summary"])):
            if field not in self.col:
                continue
            if self.ours_to_change(r, field):
                self.put(r, field, value)
            else:
                kept.append(FIELDS[field])
        self.state["emails"][email] = {"time": when.isoformat(), "key": key}
        out = {"status": "ok", "row": r, "created": created}
        if kept:
            out["note"] = f"{'、'.join(kept)}里是你自己写的内容，没有覆盖"
        return out

    def upsert(self, op: dict) -> dict:
        """Only what kol-mail sends: a channel link saved by hand, filled in when that cell is empty."""
        r = self.find(op["email"])
        if r is None:
            return {"status": "skipped", "note": "邮箱不在总表里"}
        if op.get("channel") and "channel" in self.col and not text(self.cell(r, "channel")):
            self.put(r, "channel", op["channel"])
            return {"status": "ok", "row": r, "updated": ["channel"]}
        return {"status": "ok", "row": r, "updated": []}


def norm_op(op: dict) -> dict:
    if not isinstance(op, dict) or op.get("op") not in ("event", "upsert"):
        raise UsageError(f"不认识的操作：{str(op)[:120]}")
    emails = split_emails(op.get("email"))
    if not emails:
        raise UsageError(f"{op['op']} 需要邮箱：{str(op)[:120]}")
    if op["op"] == "upsert":
        return {"op": "upsert", "email": emails[0], "channel": str(op.get("channel") or "").strip()[:300]}
    direction = {"in": "in", "收": "in", "out": "out", "发": "out"}.get(str(op.get("dir") or "").strip().lower())
    if not direction:
        raise UsageError(f"event 的 dir 只能是 in 或 out：{op.get('dir')!r}")
    subject = re.sub(r"\s+", " ", str(op.get("subject") or "")).strip()[:150]
    summary = re.sub(r"\s+", " ", str(op.get("summary") or "")).strip()[:300]
    given = bool(summary)
    if not summary:
        summary = ("对方来信：" if direction == "in" else "我方回信：") + (subject or "（无主题）")
    when = parse_when(op["time"]) if op.get("time") else now()
    when = min(when, now())  # a sender's clock in the future must not freeze the stale guard
    links = LINK.findall(str(op.get("channel_url") or ""))
    return {"op": "event", "email": emails[0], "dir": direction, "summary": summary, "summary_given": given,
            "time": when, "key": str(op.get("message_id") or "").strip()[:300],
            "name": re.sub(r"\s+", " ", str(op.get("name") or "")).strip().strip('"')[:80],
            "auto": str(op.get("type") or "").strip().endswith("自动回复"),
            "channel_url": links[0] if links else ""}


def jsonable(op: dict) -> dict:
    return {k: (v.isoformat() if isinstance(v, dt.datetime) else v) for k, v in op.items()}


def from_jsonable(op: dict) -> dict:
    return {**op, "time": parse_when(op["time"])} if op.get("op") == "event" else op


# ---------------------------------------------------------------- state, queue, locks

def load_state() -> dict:
    try:
        s = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        return s if isinstance(s, dict) else {}
    except (OSError, ValueError):
        return {}


def save_state(state: dict):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    tmp = STATE_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, STATE_FILE)


@contextlib.contextmanager
def file_lock(name: str, wait: float, what: str):
    """One byte lock on a file in STATE_DIR (msvcrt on Windows, flock on the Mac)."""
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    f = open(STATE_DIR / name, "a+")
    deadline = time.monotonic() + wait
    while True:
        try:
            f.seek(0)
            if sys.platform == "win32":
                import msvcrt
                msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            break
        except OSError:
            if time.monotonic() > deadline:
                f.close()
                raise Unavailable(what) from None
            time.sleep(0.3)
    try:
        yield
    finally:
        with contextlib.suppress(OSError):
            f.seek(0)
            if sys.platform == "win32":
                import msvcrt
                msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(f.fileno(), fcntl.LOCK_UN)
        f.close()


def read_pending() -> tuple[list[dict], int]:
    with file_lock(".pending.lock", 30, "排队文件一直被占用"):
        if not PENDING.exists():
            return [], 0
        lines = PENDING.read_text(encoding="utf-8").split("\n")
    entries = []
    for line in lines:
        if line.strip():
            with contextlib.suppress(ValueError):
                entries.append(json.loads(line))
    return entries, len(lines)


def trim_pending(done_lines: int):
    if not done_lines:
        return
    with file_lock(".pending.lock", 30, "排队文件一直被占用"):
        if not PENDING.exists():
            return
        rest = PENDING.read_text(encoding="utf-8").split("\n")[done_lines:]
        if any(x.strip() for x in rest):
            PENDING.write_text("\n".join(rest) + "\n", encoding="utf-8")
        else:
            PENDING.unlink()


def queue(ops: list[dict], source: str, reason: str):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    with file_lock(".pending.lock", 30, "排队文件一直被占用"):
        with PENDING.open("a", encoding="utf-8") as f:
            f.write(json.dumps({"queued_at": now().isoformat(), "source": source, "reason": reason,
                                "ops": [jsonable(o) for o in ops]}, ensure_ascii=False) + "\n")


def apply_ops(sheet: UserSheet, ops: list[dict]) -> list[dict]:
    out = []
    for op in ops:
        try:
            out.append(sheet.event(op) if op["op"] == "event" else sheet.upsert(op))
        except UsageError:
            raise
        except Exception as e:  # noqa: BLE001  one odd row must not stop the rest
            if is_transient(e):
                raise
            out.append({"status": "error", "error": f"{op.get('email')}: {type(e).__name__}: {e}"[:300]})
    return out


def is_transient(e: BaseException) -> bool:
    if sys.platform != "win32":
        return False
    import kol
    return kol.is_com_error(e, kol.BUSY_HRESULTS | kol.GONE_HRESULTS)


def run(ops: list[dict], source: str, mapping: dict, grid_factory=open_grid) -> dict:
    """Apply queued ops and then ops in one session, or queue ops. Exit status as in kol.py."""
    started = time.monotonic()
    deadline = started + BUSY_WAIT_SECONDS
    reason = ""
    try:
        with file_lock(".usersheet.lock", 180, "另一个写表进程一直没结束"):
            while True:
                try:
                    entries, done_lines = read_pending()
                    state = load_state()
                    with grid_factory() as grid:
                        grid.backup_once_a_day()
                        sheet = UserSheet(grid, mapping, state)
                        flushed = 0
                        for entry in entries:
                            queued_ops = [from_jsonable(o) for o in entry.get("ops", [])]
                            apply_ops(sheet, queued_ops)
                            flushed += len(queued_ops)
                        results = apply_ops(sheet, ops)
                        grid.save()
                        mode = grid.mode
                    save_state(state)
                    trim_pending(done_lines)
                    return {"status": "ok", "mode": mode, "results": results, "flushed_from_queue": flushed}
                except Unavailable as e:
                    reason = str(e)
                    deadline = min(deadline, started + UNAVAILABLE_WAIT_SECONDS)
                except UsageError:
                    raise
                except Exception as e:  # noqa: BLE001
                    if not is_transient(e):
                        raise
                    reason = "Excel 正忙（可能正在编辑单元格或开着对话框）"
                if time.monotonic() > deadline:
                    break
                time.sleep(2)
    except Unavailable as e:
        reason = str(e)
    if ops:
        queue(ops, source, reason)
    return {"status": "queued", "reason": reason, "queued": len(ops)}


# ---------------------------------------------------------------- reading without Excel

def read_file_rows(path: Path, sheet: str | None = None) -> tuple[list[str], str, list[list]]:
    g = XlsxGrid(path)
    names = g.sheet_names()
    chosen = sheet if sheet in names else None
    if chosen is None:  # the sheet with the most e-mail addresses
        best = -1
        for name in names:
            rows = g.rows(name)
            n = sum(1 for r in rows[:500] for v in r if split_emails(v))
            if n > best:
                chosen, best = name, n
    return names, chosen, g.rows(chosen)


def inspect(path: Path, sheet: str | None, header_row: int | None) -> dict:
    names, chosen, rows = read_file_rows(path, sheet)
    hr = header_row or find_header_row(rows)
    headers = rows[hr - 1] if hr - 1 < len(rows) else []
    cols = [{"col": c + 1, "header": text(h), **column_profile(rows, hr, c + 1)}
            for c, h in enumerate(headers) if text(h)]
    return {"sheets": names, "sheet": chosen, "header_row": hr, "rows": max(0, len(rows) - hr), "columns": cols,
            "guess": guess_map(rows, hr), "fields": FIELDS, "required": list(REQUIRED)}


def list_rows(mapping: dict, fields: list[str]) -> list[dict]:
    """Rows for app.py, keyed like kol.py's list (邮箱, 频道链接, 阶段, 意向); unmapped ones are empty."""
    _, _, rows = read_file_rows(WORKBOOK, mapping["sheet"])
    hr = mapping["header_row"]
    headers = rows[hr - 1] if hr - 1 < len(rows) else []
    by_text = {norm_header(h): i for i, h in reversed(list(enumerate(headers))) if text(h)}
    idx = {f: by_text.get(norm_header(h)) for f, h in mapping["columns"].items()}
    names = {"邮箱": "email", "频道链接": "channel", "KOL": "name", "待办方": "owner", "进展摘要": "summary"}
    out = []
    for r in rows[hr:]:
        item = {}
        for f in fields:
            i = idx.get(names.get(f, ""))
            item[f] = text(r[i]) if i is not None and i < len(r) else ""
        if any(item.values()):
            out.append(item)
    return out


def create(path: Path) -> dict:
    """A new, simple KOL sheet as an Excel table, already mapped."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    if path.exists():
        raise UsageError(f"这个位置已经有文件了：{path}")
    if path.suffix.lower() != ".xlsx":
        raise UsageError("新总表的文件名要以 .xlsx 结尾")
    headers = ["KOL", "平台", "频道链接", "邮箱", "最近联系时间", "待办方", "进展摘要", "备注"]
    widths = [22, 10, 42, 30, 18, 10, 60, 30]
    wb = Workbook()
    ws = wb.active
    ws.title = "KOL总表"
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E293B")
        c.alignment = Alignment(vertical="center")
    ws.append([None] * len(headers))
    ws.cell(2, 5).number_format = DEFAULT_DATE_FORMAT
    table = Table(displayName="KOL", ref=f"A1:H2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    columns = {"name": "KOL", "channel": "频道链接", "email": "邮箱", "last_contact": "最近联系时间",
               "owner": "待办方", "summary": "进展摘要"}
    return {"status": "ok", "path": str(path), "map": {"sheet": "KOL总表", "header_row": 1, "columns": columns}}


# ---------------------------------------------------------------- command line (same shape as kol.py)

def emit(obj, code: int = 0):
    print(json.dumps(obj, ensure_ascii=False, default=str))
    sys.exit(code)


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(prog="usersheet.py", description="按列对应同步同事自己的 KOL 总表")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("batch", help="执行 kol-mail 的操作：JSON 文件路径，或 - 表示从标准输入读")
    p.add_argument("file")
    p.add_argument("--source", default="")
    p.add_argument("--skip-invalid", action="store_true")
    p = sub.add_parser("list", help="读出各行（直接读文件，不启动 Excel）")
    p.add_argument("--fields", default="邮箱,频道链接")
    sub.add_parser("flush", help="把排队的改动写进总表")
    sub.add_parser("status", help="总表位置、排队数")
    p = sub.add_parser("inspect", help="读一个文件的工作表和表头，猜列对应")
    p.add_argument("--path", required=True)
    p.add_argument("--sheet")
    p.add_argument("--header-row", type=int)
    p = sub.add_parser("check", help="按列对应读一遍总表（Windows 上经 Excel），报告找到多少行")
    p = sub.add_parser("create", help="新建一张简单的 KOL 总表")
    p.add_argument("--path", required=True)
    a = ap.parse_args(argv)
    try:
        if a.cmd == "inspect":
            emit({"status": "ok", **inspect(Path(a.path), a.sheet, a.header_row)})
        if a.cmd == "create":
            emit(create(Path(a.path)))
        if a.cmd == "status":
            pending = sum(len(e.get("ops", [])) for e in read_pending()[0])
            emit({"workbook": str(WORKBOOK), "exists": WORKBOOK.is_file(), "queued_ops": pending})
        mapping = load_map()
        if a.cmd == "list":
            fields = [f.strip() for f in a.fields.split(",") if f.strip()]
            rows = list_rows(mapping, fields)
            emit({"count": len(rows), "rows": rows})
        if a.cmd == "check":
            with file_lock(".usersheet.lock", 60, "同步正在写总表，稍后再测"), open_grid(write=False) as g:
                sheet = UserSheet(g, mapping, {})
                rows = range(mapping["header_row"] + 1, sheet.last_row() + 1)
                emit({"status": "ok", "mode": g.mode, "rows": len(rows),
                      "emails": sum(1 for r in rows if split_emails(sheet.cell(r, "email"))),
                      "links": sum(1 for r in rows if LINK.search(text(sheet.cell(r, "channel")))),
                      "columns": {f: FIELDS[f] for f in sheet.col}})
        if a.cmd == "flush":
            if not read_pending()[0]:
                emit({"status": "ok", "note": "没有排队的改动"})
            res = run([], "补写队列", mapping)
        else:  # batch
            raw = sys.stdin.buffer.read() if a.file == "-" else Path(a.file).read_bytes()
            data = json.loads(raw.decode("utf-8-sig"))
            items = data.get("ops", []) if isinstance(data, dict) else data
            ops, rejected = [], []
            for item in items:
                try:
                    ops.append(norm_op(item))
                except UsageError as e:
                    if not a.skip_invalid:
                        raise
                    rejected.append({"op": item, "error": str(e)})
            res = run(ops, a.source or "kol-mail", mapping) if ops else {"status": "ok", "results": []}
            if rejected:
                res["rejected"] = rejected
        if res.get("status") != "ok":
            emit(res, 3)
        problems = res.get("rejected") or [r for r in res.get("results", []) if r.get("status") == "error"]
        emit(res, 4 if problems else 0)
    except UsageError as e:
        emit({"status": "error", "error": str(e)}, 2)
    except Unavailable as e:
        emit({"status": "unavailable", "reason": str(e)}, 3)


if __name__ == "__main__":
    main()
