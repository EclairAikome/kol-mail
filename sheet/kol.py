"""KOL 总表写入工具。

所有 AI 和程序改 KOL总表.xlsx 都只走这里，不要用 openpyxl 之类的库直接改文件。

总表在 Excel 里开着时，通过 COM 直接写进打开的窗口并保存，主人当场就能看到；没开时在后台起一个
看不见的 Excel 写完再关。Excel 正忙（主人正在单元格里打字、开着对话框）时会等着重试，等不到就把
改动排进 pending.jsonl，之后任何一次写入调用都会先把排队的补写进去。

Mac 版没有 Excel COM，总表是一张 Google 表格：KOL_SHEET_PATH 写成 gsheet:<表格ID> 时，读写都走 gsheet.py，
下面的记账规则不变。

用法见同目录 README.md，或运行 `python kol.py -h`。
"""

from __future__ import annotations

import argparse
import contextlib
import datetime as dt
import email.utils
import gc
import hashlib
import json
import os
import re
import sys
import time
from pathlib import Path

if sys.platform == "win32":
    import msvcrt
else:  # the Mac build (Google Sheets backend) only needs the lock
    import fcntl

HERE = Path(__file__).resolve().parent
WORKBOOK = Path(os.environ.get("KOL_SHEET_PATH") or HERE.parent / "KOL总表.xlsx")
STATE_DIR = Path(os.environ.get("KOL_SHEET_STATE") or HERE)
_SETTING = (os.environ.get("KOL_SHEET_PATH") or "").strip()
GSHEET_ID = _SETTING[len("gsheet:"):].strip() or None if _SETTING.startswith("gsheet:") else None  # the Mac build
PENDING = STATE_DIR / "pending.jsonl"
FAILED = STATE_DIR / "pending.failed.jsonl"
LOCKFILE = STATE_DIR / ".kol.lock"
BACKUPS = STATE_DIR / "backups"
KEEP_BACKUPS = 14
DEFAULT_PROJECT = os.environ.get("KOL_DEFAULT_PROJECT", "")

BUSY_WAIT_SECONDS = 60     # how long one command keeps retrying a busy Excel before queueing
UNAVAILABLE_WAIT_SECONDS = 8  # read-only / locked by another program: retry briefly, then queue
LOCK_WAIT_SECONDS = 180    # how long to wait for another kol.py process to finish

MAIN_SHEET, LOG_SHEET, BATCH_SHEET, DICT_SHEET = "KOL总表", "沟通记录", "发信批次", "说明"
MAIN_TABLE, LOG_TABLE, BATCH_TABLE, PROJECT_TABLE = "KOL", "LOG", "BATCH", "PROJ"

# Added on 2026-10-29: the Telegram / Discord accounts of a KOL, so a chat finds its row the way an address does.
# Sheets made before have no such column and keep working: locating by chat finds nothing there, and writing a chat
# account only warns (CHAT_MIGRATE_HINT).
CHAT_HEADER = "聊天账号"
OPTIONAL_HEADERS = [CHAT_HEADER]
CHAT_MIGRATE_HINT = f"在“{MAIN_SHEET}”页表格最右边紧挨着加一列，表头写“{CHAT_HEADER}”"
MAIN_HEADERS = ["ID", "KOL", "项目", "语言", "阶段", "意向", "待办方", "下一步",
                "对方报价", "方案/条件", "频道数据", "进展摘要", "备注", "平台", "频道", "邮箱", "首次发信",
                "更新时间", "更新来源", "频道链接", CHAT_HEADER]
LOG_HEADERS = ["时间", "ID", "KOL", "方向", "类型", "摘要", "渠道", "邮箱", "来源", "键"]
BATCH_HEADERS = ["批次", "项目", "语言", "日期", "人数", "主题", "Message-ID"]
FORMULA_COLUMNS = {"频道"}  # calculated column, never written by this tool
STALE_TOLERANCE = dt.timedelta(seconds=90)  # a Date header vs our own send stamp, plus HH:MM truncation of "last"

STAGES = ["未联系", "已发信", "洽谈中", "已谈成", "已婉拒", "退信"]
INTENTS = ["待定", "纯固薪", "CPS+固薪", "纯CPS"]
OWNERS = ["我方", "对方"]
DIRECTIONS = {"收": "收", "in": "收", "incoming": "收", "发": "发", "out": "发", "outgoing": "发"}

# field name (CLI / JSON) -> column header. Chinese headers are accepted as field names too.
FIELDS = {
    "id": "ID", "kol": "KOL", "project": "项目", "lang": "语言", "stage": "阶段", "intent": "意向",
    "owner": "待办方", "next": "下一步", "quote": "对方报价", "terms": "方案/条件",
    "stats": "频道数据", "summary": "进展摘要", "note": "备注", "platform": "平台", "email": "邮箱",
    "first": "首次发信", "channel": "频道链接", "chat": CHAT_HEADER,
}
HEADER_TO_FIELD = {h: f for f, h in FIELDS.items()} | {"频道": "channel"}
# "last" (双方最近一封邮件的时间) is not a column. It only tells kol.py how fresh the caller's view of the
# conversation is, so an older judgement never overwrites state that kol-mail recorded after it.
GUARD_KEYS = {"last", "最近往来"}
DATE_FIELDS = {"first"}
TEXT_LIMITS = {"summary": 300, "next": 200, "quote": 300, "terms": 400, "stats": 300, "note": 1000}

EPOCH = dt.datetime(1899, 12, 30)
INVISIBLE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f­​-‏‪-‮⁠-⁩﻿]")

# COM failures that mean "Excel is busy right now, try again later".
BUSY_HRESULTS = {
    -2147418111,  # RPC_E_CALL_REJECTED: Excel is in cell-edit mode or showing a dialog
    -2147417846,  # RPC_E_SERVERCALL_RETRYLATER
    -2146777998,  # 0x800AC472 VBA_E_IGNORE: Excel refused the call while in a modal state
}
GONE_HRESULTS = {
    -2147023174,  # RPC_S_SERVER_UNAVAILABLE: Excel closed while we were talking to it
    -2147417848,  # RPC_E_DISCONNECTED
    -2147023170,  # RPC_S_CALL_FAILED
}


class UsageError(Exception):
    """Bad input. Nothing has been written."""


class Unavailable(Exception):
    """The workbook can't be written right now; the operations get queued."""


# ---------------------------------------------------------------- small helpers

def now() -> dt.datetime:
    return dt.datetime.now().replace(microsecond=0)


def to_serial(d: dt.datetime) -> float:
    return round((d - EPOCH).total_seconds() / 86400, 8)


def from_serial(x) -> dt.datetime | None:
    if x in (None, ""):
        return None
    if isinstance(x, dt.datetime):
        return x.replace(tzinfo=None)
    return EPOCH + dt.timedelta(seconds=round(float(x) * 86400))


def parse_when(value) -> dt.datetime:
    """ISO dates/times, 2026/10/16, or an e-mail Date header. Time zones are converted to local time."""
    if isinstance(value, dt.datetime):
        d = value
    else:
        s = str(value).strip().replace("/", "-")
        try:
            d = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
        except ValueError:
            try:
                d = email.utils.parsedate_to_datetime(str(value))
            except (TypeError, ValueError, IndexError):
                raise UsageError(f"看不懂的时间：{value!r}，请用 2026-10-16 或 2026-10-16 14:30") from None
    if d.tzinfo is not None:
        d = d.astimezone().replace(tzinfo=None)
    return d.replace(microsecond=0)


def fmt_when(d: dt.datetime | None) -> str:
    if d is None:
        return ""
    return d.strftime("%Y-%m-%d") if (d.hour, d.minute, d.second) == (0, 0, 0) else d.strftime("%Y-%m-%d %H:%M")


def clean(value, limit: int = 600) -> str:
    """Single-line plain text: invisible characters removed, line breaks folded into ' / '."""
    s = str(value).replace("\u2028", "\n").replace("\u2029", "\n").replace("\x85", "\n")
    s = INVISIBLE.sub("", s).replace("\r\n", "\n").replace("\r", "\n").strip()
    s = re.sub(r"\s*\n\s*", " / ", s)
    s = re.sub(r"[ \t]{2,}", " ", s)
    return s[:limit]


def text_cell(s: str) -> str:
    """A leading apostrophe makes Excel store the value as text, so '=...' from an e-mail can never become a formula."""
    return "'" + s if s else ""


def split_emails(value) -> list[str]:
    return [e.strip("<>()（）") for e in re.split(r"[;,，；、/\s]+", str(value or "").lower()) if "@" in e]


# A chat account is "tg:<Telegram user id>" or "dc:<Discord user id>", optionally followed by "(<display name>)";
# a 聊天账号 cell holds several separated by "; ". Only prefix and digits identify the person: display names change.
CHAT_TOKEN = re.compile(r"\s*(tg|dc)\s*:\s*(\d{1,20})\s*(?:\((.*)\))?\s*", re.IGNORECASE | re.DOTALL)
CHAT_ID_IN_TEXT = re.compile(r"(?<![a-z0-9])(tg|dc)\s*:\s*(\d{1,20})(?!\d)", re.IGNORECASE)
CHAT_PLATFORMS = {"tg": "Telegram", "dc": "Discord"}
CHAT_EXAMPLE = "tg:1274797389(@yummytim) 或 dc:515534399828131853(ShiruTV_)"


def parse_chat(piece) -> tuple[str, str] | None:
    """(identity, display name) of one token, e.g. ("dc:515534399828131853", "ShiruTV_"); None if it is not one."""
    m = CHAT_TOKEN.fullmatch(str(piece or ""))
    return (f"{m.group(1).lower()}:{m.group(2)}", (m.group(3) or "").strip()) if m else None


def chat_pieces(cell) -> list[str]:
    return [p.strip() for p in re.split(r"[;；]", str(cell or "")) if p.strip()]


def chat_ids(cell) -> list[str]:
    """Identities in a 聊天账号 cell. A piece that is not a clean token (typed by hand, "tg:1, dc:2") still counts for
    the ids in it; display names never do, so a name like "dc:1" can't make a match."""
    out = []
    for piece in chat_pieces(cell):
        token = parse_chat(piece)
        found = [token[0]] if token else [f"{a.lower()}:{b}" for a, b in CHAT_ID_IN_TEXT.findall(piece)]
        out += [x for x in found if x not in out]
    return out


def chat_token(ident: str, display: str = "") -> str:
    return f"{ident}({display})" if display else ident


def chat_platform(token: str) -> str:
    return CHAT_PLATFORMS.get(str(token)[:2].lower(), "")


def norm_chats(value) -> list[str]:
    """An op's chat value (one token, several joined by ';', or a list) as clean tokens, one per identity."""
    found: dict[str, str] = {}
    for item in (value if isinstance(value, list) else [value]):
        for piece in chat_pieces(item):
            token = parse_chat(clean(piece, 300))
            if token is None:
                raise UsageError(f"聊天账号格式不对：{piece[:80]!r}，应该像 {CHAT_EXAMPLE}")
            ident, display = token
            found[ident] = clean(display, 60) or found.get(ident, "")
    return [chat_token(i, d) for i, d in found.items()]


def merge_chats(cell, tokens: list[str]) -> str:
    """The cell with tokens merged in, like addresses: a new account is appended, a known one gets the newer display
    name, nothing is ever dropped. Pieces typed by hand stay as typed; an unchanged cell comes back unchanged."""
    old = "" if cell is None else str(cell)
    pieces = chat_pieces(old)
    changed = False
    for token in tokens:
        ident, display = parse_chat(token)
        for k, piece in enumerate(pieces):
            have = parse_chat(piece)
            if have and have[0] == ident:
                if display and have[1] != display:
                    pieces[k], changed = chat_token(ident, display), True
                break
            if not have and ident in chat_ids(piece):
                break
        else:
            pieces.append(chat_token(ident, display))
            changed = True
    return "; ".join(pieces) if changed else old


def clean_url(value) -> str:
    u = clean(value, 500)
    if u and not re.fullmatch(r"https?://[^\s\"'<>]+", u):
        raise UsageError(f"频道链接必须是 http(s) 网址：{u[:80]}")
    return u


def platform_of(url: str) -> str:
    host = re.sub(r"^https?://(www\.|m\.)?", "", url.lower()).split("/")[0]
    for key, name in (("youtube", "YouTube"), ("youtu.be", "YouTube"), ("twitch", "Twitch"), ("tiktok", "TikTok"),
                      ("instagram", "Instagram"), ("bilibili", "Bilibili"), ("x.com", "X"), ("twitter", "X")):
        if key in host:
            return name
    return ""


def is_com_error(exc: BaseException, codes: set[int]) -> bool:
    try:
        import pywintypes
    except ImportError:
        return False
    if not isinstance(exc, pywintypes.com_error):
        return False
    found = {exc.hresult}
    if exc.excepinfo and len(exc.excepinfo) > 5 and exc.excepinfo[5]:
        found.add(exc.excepinfo[5])
    return bool(found & codes)


# ---------------------------------------------------------------- operation normalisation (before touching Excel)

def opt(op: dict, key: str, default=""):
    """op[key] with None treated as missing (the CLI passes None for every option left out)."""
    v = op.get(key)
    return default if v is None else v


def as_bool(v, default: bool) -> bool:
    if v is None:
        return default
    if isinstance(v, str):
        return v.strip().lower() not in ("", "false", "0", "no", "n", "否")
    return bool(v)


def not_future(d: dt.datetime, what: str) -> dt.datetime:
    if d > now() + dt.timedelta(days=1):
        raise UsageError(f"{what}不能晚于明天：{d:%Y-%m-%d %H:%M}（年份写错了？）")
    return d


def norm_choice(value, allowed: list[str], what: str, blank_ok: bool = True) -> str:
    v = clean(value, 20)
    if v in ("", "-", "无", "none") and blank_ok:
        return ""
    if v not in allowed:
        raise UsageError(f"{what}只能是 {' / '.join(allowed)}，收到 {value!r}")
    return v


def norm_fields(raw: dict) -> dict:
    """Validate upsert fields. Returns {field: value}; '' means clear the cell."""
    out = {}
    for key, value in raw.items():
        if key in ("op", "source", "create", "rename", "only_if_stage"):
            continue
        if key in GUARD_KEYS:
            if value is not None and str(value).strip():
                out["last"] = not_future(parse_when(value), "last")
            continue
        field = key if key in FIELDS else HEADER_TO_FIELD.get(key)
        if field is None:
            raise UsageError(f"不认识的字段：{key}。可用字段：{', '.join(FIELDS)}")
        if value is None:
            continue
        if field == "stage":
            out[field] = norm_choice(value, STAGES, "阶段", blank_ok=False)
        elif field == "intent":
            out[field] = norm_choice(value, INTENTS, "意向")
        elif field == "owner":
            out[field] = norm_choice(value, OWNERS, "待办方")
        elif field in DATE_FIELDS:
            out[field] = not_future(parse_when(value), "日期") if str(value).strip() else ""
        elif field == "email":
            emails = split_emails(value)
            if str(value).strip() and not emails:
                raise UsageError(f"邮箱格式不对：{value!r}")
            out[field] = emails
        elif field == "chat":
            out[field] = norm_chats(value)
        elif field == "channel":
            out[field] = clean_url(value)
        elif field == "lang":
            out[field] = clean(value, 12).upper()
        else:
            out[field] = clean(value, TEXT_LIMITS.get(field, 120))
    return out


def norm_op(op: dict) -> dict:
    if not isinstance(op, dict) or "op" not in op:
        raise UsageError(f"每条操作都要有 op 字段：{str(op)[:120]}")
    kind = op["op"]
    if kind == "upsert":
        fields = norm_fields(op)
        if not ({"id", "email", "kol", "channel", "chat"} & fields.keys()):
            raise UsageError("upsert 至少要给 id、email、chat、channel、kol 其中一个来定位 KOL")
        only = op.get("only_if_stage")
        if only is not None:
            only = [norm_choice(x, STAGES, "only_if_stage 里的阶段", blank_ok=False)
                    for x in (only if isinstance(only, list) else [only])]
        return {"op": "upsert", "fields": fields, "create": as_bool(op.get("create"), True),
                "rename": as_bool(op.get("rename"), False), "only_if_stage": only,
                "source": clean(opt(op, "source"), 60)}
    if kind in ("log", "event"):
        direction = DIRECTIONS.get(str(op.get("dir", "")).strip().lower())
        if direction is None:
            raise UsageError(f"{kind} 的 dir 只能是 收/发（或 in/out），收到 {op.get('dir')!r}")
        who = norm_fields({k: op[k] for k in ("id", "email", "kol", "chat") if op.get(k)})
        if not who:
            raise UsageError(f"{kind} 至少要给 id、email、chat、kol 其中一个")
        # In log and event, channel is 渠道 (邮件 / Telegram / Discord), not 频道链接 as in upsert. A chat message
        # without one gets its platform; everything else stays 邮件, as callers that send none expect.
        channel = clean(opt(op, "channel"), 300)
        if re.match(r"(?i)(https?://|www\.)", channel):
            raise UsageError(f"{kind} 的 channel 是渠道（邮件、Telegram、Discord），频道链接请用 upsert 写：{channel[:80]}")
        platform = chat_platform(who["chat"][0]) if who.get("chat") else ""
        subject = clean(opt(op, "subject"), 200)
        summary = clean(opt(op, "summary"), 400)
        etype = clean(opt(op, "type"), 12) or platform or "回信"
        if not summary:
            if not subject:
                raise UsageError(f"{kind} 需要 summary 或 subject")
            if platform:
                summary = ("对方消息：" if direction == "收" else "我方消息：") + subject
            else:
                summary = ("收到回信：" if direction == "收" else "我方回信：") + subject
        when = parse_when(op["time"]) if op.get("time") else now()
        if when > now() + STALE_TOLERANCE:
            when = now()  # a sender's clock in the future must not freeze the stale-update guard
        return {"op": kind, "who": who, "dir": direction, "type": etype, "summary": summary, "time": when,
                "channel": channel[:12] or platform or "邮件", "source": clean(opt(op, "source"), 60),
                "key": clean(op.get("message_id") or op.get("key") or "", 300)}
    if kind == "add_batch":
        if not op.get("batch") or not op.get("subject"):
            raise UsageError("add_batch 需要 batch 和 subject")
        digits = re.sub(r"\D", "", str(opt(op, "count")))
        out = {"op": "add_batch", "batch": clean(op["batch"], 40), "project": clean(op.get("project") or DEFAULT_PROJECT, 40),
               "lang": clean(opt(op, "lang"), 12).upper(), "date": not_future(parse_when(op.get("date") or now()), "日期"),
               "count": int(digits) if digits else 0, "subject": clean(op["subject"], 300),
               "key": clean(opt(op, "message_id"), 300)}
        if not out["key"]:  # retries must not add the same batch twice
            raw = f"{out['batch']}|{out['lang']}|{out['date']:%Y-%m-%d}|{out['subject']}"
            out["key"] = "auto:" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]
        return out
    raise UsageError(f"不认识的操作：{kind}（可用 upsert / log / event / add_batch）")


def to_jsonable(op: dict) -> dict:
    def conv(v):
        if isinstance(v, dt.datetime):
            return {"$dt": v.isoformat()}
        if isinstance(v, dict):
            return {k: conv(x) for k, x in v.items()}
        if isinstance(v, list):
            return [conv(x) for x in v]
        return v
    return conv(op)


def from_jsonable(op: dict) -> dict:
    def conv(v):
        if isinstance(v, dict):
            if set(v) == {"$dt"}:
                return dt.datetime.fromisoformat(v["$dt"])
            return {k: conv(x) for k, x in v.items()}
        if isinstance(v, list):
            return [conv(x) for x in v]
        return v
    return conv(op)


# ---------------------------------------------------------------- Excel connection

def find_open_workbook(path: Path):
    """The workbook object if this file is open in any Excel instance on this machine (via the Running Object Table)."""
    import pythoncom
    import win32com.client
    target = os.path.normcase(str(path.resolve()))
    rot = pythoncom.GetRunningObjectTable()
    ctx = pythoncom.CreateBindCtx(0)
    for moniker in rot:
        try:
            name = moniker.GetDisplayName(ctx, None)
        except Exception:
            continue
        if os.path.normcase(name) == target:
            obj = rot.GetObject(moniker)
            return win32com.client.Dispatch(obj.QueryInterface(pythoncom.IID_IDispatch))
    return None


def file_locked(path: Path) -> bool:
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True


def process_of(app) -> int | None:
    with contextlib.suppress(Exception):
        import win32process
        return win32process.GetWindowThreadProcessId(app.Hwnd)[1]
    return None


def make_sure_exited(pid: int | None, wait: float = 8.0):
    """A headless Excel that hangs around keeps the file locked and every later call would queue. Never leave one."""
    if not pid:
        return
    import win32api
    import win32con
    import win32event
    try:
        handle = win32api.OpenProcess(win32con.SYNCHRONIZE | win32con.PROCESS_TERMINATE, False, pid)
    except Exception:
        return  # already gone
    try:
        if win32event.WaitForSingleObject(handle, int(wait * 1000)) != win32event.WAIT_OBJECT_0:
            win32api.TerminateProcess(handle, 1)
    finally:
        win32api.CloseHandle(handle)


def open_book(write: bool = True):
    """The master workbook for one command: the Google sheet in the Mac build, else Excel."""
    if GSHEET_ID:
        import gsheet
        return gsheet.GBook(GSHEET_ID, STATE_DIR, write=write, usage=UsageError, unavailable=Unavailable,
                            formula_columns=FORMULA_COLUMNS)
    return Book(write=write)


class Book:
    """One connection to the master workbook for the duration of a command."""

    def __init__(self, path: Path = WORKBOOK, write: bool = True):
        self.path, self.write = path, write
        self.wb = self.app = None
        self.mode = None
        self.pid = None

    def __enter__(self):
        import pythoncom
        pythoncom.CoInitialize()
        try:
            return self._connect()
        except BaseException:
            self.__exit__(None, None, None)  # never leave a half-started Excel behind
            raise

    def _connect(self):
        import win32com.client
        if not self.path.exists():
            raise UsageError(f"找不到总表：{self.path}")
        self.wb = find_open_workbook(self.path)
        if self.wb is not None:
            app = self.wb.Application
            # An invisible instance nobody controls is a leftover of an earlier crashed run: use it, then close it.
            self.mode = "live" if (app.Visible or app.UserControl) else "orphan"
            if self.mode == "orphan":
                self.app, self.pid = app, process_of(app)
        else:
            if file_locked(self.path):
                raise Unavailable("总表被别的程序占用了（不在本机 Excel 的已打开列表里，可能是受保护视图或别的软件）")
            self.app = win32com.client.DispatchEx("Excel.Application")
            self.pid = process_of(self.app)
            self.app.Visible = False
            self.app.DisplayAlerts = False
            self.app.ScreenUpdating = False
            self.app.EnableEvents = False
            self.app.AskToUpdateLinks = False
            self.wb = self.app.Workbooks.Open(str(self.path), 0, False)  # UpdateLinks=0, ReadOnly=False
            self.mode = "headless"
        if self.write and self.wb.ReadOnly:
            raise Unavailable("总表在 Excel 里是以只读方式打开的，存不进去。请关掉后正常打开")
        return self

    def __exit__(self, exc_type, exc, tb):
        import pythoncom
        ours = self.mode in ("headless", "orphan") or (self.mode is None and self.app is not None)
        if ours and self.mode == "orphan":
            with contextlib.suppress(Exception):
                if self.app.Workbooks.Count > 1:  # the hidden instance holds other files too: only close ours
                    self.wb.Close(False)
                    ours = False
        try:
            if ours:
                with contextlib.suppress(Exception):
                    if self.wb is not None:
                        self.wb.Close(False)
                with contextlib.suppress(Exception):
                    self.app.Quit()
        finally:
            self.wb = self.app = None
            gc.collect()
            if ours:
                make_sure_exited(self.pid)
            pythoncom.CoUninitialize()

    def table(self, sheet: str, name: str) -> "Table":
        try:
            return Table(self.wb.Worksheets(sheet).ListObjects(name))
        except Exception as e:
            raise UsageError(f"总表里找不到 {sheet} 页的表格 {name}（被改名或删掉了？）：{e}") from None

    def save(self):
        self.wb.Save()

    def backup_once_a_day(self):
        BACKUPS.mkdir(parents=True, exist_ok=True)
        target = BACKUPS / f"{self.path.stem}_{dt.date.today():%Y%m%d}.xlsx"
        if not target.exists():
            self.wb.SaveCopyAs(str(target))
        old = sorted(BACKUPS.glob(f"{self.path.stem}_*.xlsx"))[:-KEEP_BACKUPS]
        for f in old:
            with contextlib.suppress(OSError):
                f.unlink()


class Table:
    def __init__(self, lo):
        self.lo = lo
        self.headers = [str(h) if h is not None else "" for h in lo.HeaderRowRange.Value2[0]]
        self.col = {h: i + 1 for i, h in enumerate(self.headers)}

    def need(self, headers: list[str]):
        missing = [h for h in headers if h not in self.col]
        if missing:
            raise UsageError(f"表格 {self.lo.Name} 缺少列：{', '.join(missing)}（列名被改过？改回去再试）")

    def values(self) -> list[list]:
        body = self.lo.DataBodyRange
        if body is None:
            return []
        v = body.Value2
        if not isinstance(v, tuple):
            return [[v]]
        return [list(r) for r in v]

    def get(self, row: list, header: str):
        return row[self.col[header] - 1]

    def cell(self, r: int, header: str):
        """The live value of one cell (row r of the body)."""
        return self.lo.DataBodyRange.Cells(r, self.col[header]).Value2

    def set(self, r: int, header: str, value, number_format: str | None = None):
        cell = self.lo.DataBodyRange.Cells(r, self.col[header])
        if number_format:
            cell.NumberFormat = number_format
        cell.Value2 = value

    def write_row(self, r: int, values: dict):
        """Fill row r in a single COM call. Formula cells keep their formulas (read back through .Formula)."""
        rng = self.lo.ListRows(r).Range
        current = list(rng.Formula[0])
        for h, v in values.items():
            current[self.col[h] - 1] = v
        rng.Formula = (tuple(current),)

    def blank(self, row: list) -> bool:
        return all(v in (None, "") for h, v in zip(self.headers, row) if h not in FORMULA_COLUMNS)

    def append(self) -> int:
        """Index (1-based, within the body) of a fresh row at the bottom. A blank last row left by an interrupted
        run is reused. Always the bottom: Excel refuses to insert rows inside a table while a filter is on."""
        rows = self.values()
        if rows and self.blank(rows[-1]):
            return len(rows)
        self.lo.ListRows.Add()
        return self.lo.ListRows.Count


# ---------------------------------------------------------------- the three tables

class Master:
    def __init__(self, book: Book):
        self.book = book
        self.t = book.table(MAIN_SHEET, MAIN_TABLE)
        self.t.need([h for h in MAIN_HEADERS if h not in FORMULA_COLUMNS and h not in OPTIONAL_HEADERS])
        self.rows = self.t.values()

    @property
    def has_chat(self) -> bool:
        return CHAT_HEADER in self.t.col

    def chat_warning(self, fields: dict, what: str) -> str:
        """Why the chat account in `fields` did nothing on a sheet without the column; "" when it has it."""
        if not fields.get("chat") or self.has_chat:
            return ""
        return f"总表还没有“{CHAT_HEADER}”列，{what}。{CHAT_MIGRATE_HINT}"

    def refresh(self):
        self.rows = self.t.values()

    def find(self, who: dict) -> int | None:
        g = self.t.get
        if who.get("id"):
            for i, row in enumerate(self.rows, 1):
                if str(g(row, "ID") or "").strip().upper() == who["id"].upper():
                    return i
        for address in who.get("email") or []:
            for i, row in enumerate(self.rows, 1):
                if address in split_emails(g(row, "邮箱")):
                    return i
        if who.get("chat") and self.has_chat:
            want = {parse_chat(t)[0] for t in who["chat"]}
            for i, row in enumerate(self.rows, 1):
                if want & set(chat_ids(g(row, CHAT_HEADER))):
                    return i
        if who.get("channel"):
            want = who["channel"].rstrip("/").lower()
            for i, row in enumerate(self.rows, 1):
                if str(g(row, "频道链接") or "").rstrip("/").lower() == want:
                    return i
        if who.get("kol") and not (who.get("email") or who.get("chat")):
            # A name only counts when no address or chat account was given: anyone can write in under a known
            # channel's name, and matching them by name would merge a stranger's address into the real KOL's row.
            hits = self.name_hits(who["kol"])
            if len(hits) > 1:
                raise UsageError(f"名叫 {who['kol']} 的 KOL 有 {len(hits)} 个，请改用 id 或 email 定位")
            if hits:
                return hits[0]
        return None

    def name_hits(self, name: str) -> list[int]:
        return [i for i, row in enumerate(self.rows, 1)
                if str(self.t.get(row, "KOL") or "").strip().casefold() == name.strip().casefold()]

    def record(self, i: int) -> dict:
        row = self.rows[i - 1]
        out = {}
        for h in self.t.headers:
            v = self.t.get(row, h)
            if h in ("首次发信", "更新时间"):
                v = fmt_when(from_serial(v)) if isinstance(v, (int, float)) and v else (v or "")
            elif isinstance(v, float) and v.is_integer():
                v = int(v)
            out[h] = "" if v is None else v
        return out

    def latest_contact(self, i: int) -> dt.datetime | None:
        """When we last heard from or wrote to the KOL in row i, taken from 沟通记录. Log rows are matched by ID, or
        by address when they carry no ID; automatic replies and times ahead of our clock don't count."""
        row = self.rows[i - 1]
        kid = str(self.t.get(row, "ID") or "").strip()
        emails = set(split_emails(self.t.get(row, "邮箱")))
        t = self.book.table(LOG_SHEET, LOG_TABLE)
        ceiling = now() + STALE_TOLERANCE
        best = None
        for r in t.values():
            rid = str(t.get(r, "ID") or "").strip()
            mine = (kid and rid == kid) or (not rid and emails and set(split_emails(t.get(r, "邮箱"))) & emails)
            if not mine or is_auto_reply(t.get(r, "类型")):
                continue
            v = t.get(r, "时间")
            if isinstance(v, (int, float)) and v:
                d = from_serial(v)
                if d <= ceiling and (best is None or d > best):
                    best = d
        return best

    def next_id(self) -> str:
        nums = [int(m.group(1)) for row in self.rows if (m := re.fullmatch(r"K(\d+)", str(self.t.get(row, "ID") or "")))]
        return f"K{(max(nums) + 1 if nums else 1):04d}"

    def write(self, i: int, header: str, value):
        """Write one cell and keep the cached row in step."""
        fmt = None
        if isinstance(value, dt.datetime):
            cell, fmt = to_serial(value), ("yyyy-mm-dd hh:mm" if header == "更新时间" else "yyyy-mm-dd")
        elif isinstance(value, str):
            cell = text_cell(value)
        else:
            cell = value
        self.t.set(i, header, cell, fmt)
        self.rows[i - 1][self.t.col[header] - 1] = to_serial(value) if isinstance(value, dt.datetime) else value

    def upsert(self, fields: dict, source: str, create: bool = True, rename: bool = False,
               only_if_stage: list[str] | None = None) -> dict:
        warning = self.chat_warning(fields, "聊天账号没有写进去，其余字段照常写入")
        out = self._upsert(fields, source, create, rename, only_if_stage)
        if warning:
            out["warning"] = warning
        return out

    def _upsert(self, fields: dict, source: str, create: bool, rename: bool, only_if_stage: list[str] | None) -> dict:
        i = self.find(fields)  # with every locator, so a chat account still keeps the name from matching
        who = {k: fields.get(k) for k in ("id", "email", "chat", "kol") if fields.get(k)}
        if not self.has_chat:
            fields = {k: v for k, v in fields.items() if k != "chat"}
        if i is not None:
            stage = self.t.get(self.rows[i - 1], "阶段") or ""
            if only_if_stage and stage not in only_if_stage:
                rec = self.record(i)
                return {"status": "skipped_stage", "id": rec["ID"], "kol": rec["KOL"], "stage": stage}
            return self.update(i, fields, source, rename=rename)
        if not create:
            return {"status": "not_found", "who": who}
        known = who.get("email") or who.get("chat")
        if known and fields.get("kol") and self.name_hits(fields["kol"]):
            rec = self.record(self.name_hits(fields["kol"])[0])
            return {"status": "needs_review", "who": who,
                    "error": f"{', '.join(known)} 不在总表里，但已有同名 KOL {rec['KOL']}（{rec['ID']}）。"
                             f"确认是同一个人就带上 \"id\": \"{rec['ID']}\" 再写；不是就换个名字"}
        if not (fields.get("kol") or fields.get("email")):
            raise UsageError("新建 KOL 至少要给 kol 名称或 email")
        i = self.t.append()
        self.refresh()
        fields = dict(fields)
        defaults = {"id": self.next_id(), "kol": fields["email"][0].split("@")[0] if fields.get("email") else "",
                    "project": DEFAULT_PROJECT, "stage": "已发信" if fields.get("first") else "未联系",
                    "platform": platform_of(fields.get("channel") or "")}
        for k, v in defaults.items():
            if not fields.get(k) and v:
                fields[k] = v
        values = {}
        for field, value in fields.items():
            if field == "last":
                continue
            if field == "email":
                value = "; ".join(value)
            elif field == "chat":
                value = merge_chats("", value)
            elif field == "note" and value:
                value = f"[{now():%m-%d}] {value}"
            values[FIELDS[field]] = to_serial(value) if isinstance(value, dt.datetime) else text_cell(value)
        values["更新时间"] = to_serial(now())
        values["更新来源"] = text_cell(source or "未注明")
        self.t.write_row(i, values)
        self.refresh()
        if fields.get("project"):
            ensure_project(self.book, fields["project"])
        rec = self.record(i)
        return {"status": "created", "id": rec["ID"], "kol": rec["KOL"],
                "changed": [h for h in values if h not in ("更新时间", "更新来源")]}

    def update(self, i: int, fields: dict, source: str, created: bool = False, rename: bool = False) -> dict:
        expect = self.t.get(self.rows[i - 1], "ID")
        actual = self.t.cell(i, "ID")
        if expect and actual != expect:  # rows moved since we read them
            self.refresh()
            j = self.find({"id": str(expect)})
            if not j:
                raise UsageError(f"{expect} 这一行在写入途中被删除了")
            i = j
        row = self.rows[i - 1]
        changed, kept = [], []
        fields = dict(fields)
        if not self.has_chat:
            fields.pop("chat", None)  # nowhere to put it; upsert() said so in its warning
        new_last = fields.pop("last", None)
        cur_last = self.latest_contact(i) if new_last and not created else None
        if new_last and cur_last:
            # A date-only "last" counts as 00:00 of that day: anything logged later that day wins.
            if new_last < cur_last - STALE_TOLERANCE:
                # Someone (kol-mail) recorded a newer mail after this caller looked at the conversation: keep the
                # newer 待办方 / 下一步 / 阶段 rather than overwrite them with a stale judgement.
                kept = [FIELDS[f] for f in ("owner", "next", "stage") if f in fields]
                fields = {k: v for k, v in fields.items() if k not in ("owner", "next", "stage")}
        for field, value in fields.items():
            header = FIELDS[field]
            old = self.t.get(row, header)
            if field == "id" and not created:
                continue
            if field == "kol" and not (created or rename):
                continue  # the name only locates the row; renaming needs rename=true
            if field == "email":
                known = split_emails(old)
                new = [e for e in value if e not in known]
                if not new:
                    continue
                value = (str(old).rstrip("; ") + "; " if old else "") + "; ".join(new)
            elif field == "chat":
                value = merge_chats(old, value)  # unchanged cell -> equal to old, skipped below
            elif field == "note":
                if not value:
                    continue
                stamp = f"[{now():%m-%d}] {value}"
                if old and stamp in str(old):
                    continue  # already appended (a retry after an interruption)
                value = f"{old}；{stamp}" if old else stamp
            elif field == "first" and value:
                old_d = from_serial(old) if isinstance(old, (int, float)) else None
                if old_d and old_d <= value:
                    continue  # 首次发信 only moves back
            if (old if old is not None else "") == value or (isinstance(value, dt.datetime) and old == to_serial(value)):
                continue
            self.write(i, header, value)
            changed.append(header)
        if changed or created:
            self.write(i, "更新时间", now())
            self.write(i, "更新来源", source or "未注明")
        if fields.get("project"):
            ensure_project(self.book, fields["project"])
        rec = self.record(i)
        out = {"status": "created" if created else ("updated" if changed else "unchanged"),
               "id": rec["ID"], "kol": rec["KOL"], "changed": changed}
        if kept:
            out["kept_newer"] = kept
        return out


class Log:
    def __init__(self, book: Book):
        self.t = book.table(LOG_SHEET, LOG_TABLE)
        self.t.need(LOG_HEADERS)

    @staticmethod
    def key_for(op: dict) -> str:
        if op.get("key"):
            return op["key"]
        # chat identities, not tokens: a changed display name must not make the same message a new row
        chats = ",".join(parse_chat(t)[0] for t in op["who"].get("chat") or [])
        who = ",".join(op["who"].get("email") or []) or op["who"].get("id") or chats or op["who"].get("kol", "")
        raw = f"{op['time'].isoformat()}|{who}|{op['dir']}|{op['type']}|{op['summary']}"
        return "auto:" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]

    def find(self, key: str) -> int | None:
        for i, row in enumerate(self.t.values(), 1):
            if str(self.t.get(row, "键") or "") == key:
                return i
        return None

    def add(self, op: dict, master: Master, source: str) -> dict:
        key = self.key_for(op)
        i = self.find(key)
        if i:
            row = self.t.values()[i - 1]
            if op["summary"] and op.get("update_summary") and self.t.get(row, "摘要") != op["summary"]:
                self.t.set(i, "摘要", text_cell(op["summary"]))
                return {"status": "summary_updated", "key": key}
            return {"status": "exists", "key": key}
        i_kol = master.find(op["who"])
        if i_kol:
            rec = master.record(i_kol)
            kid, name, mail = rec["ID"], rec["KOL"], (op["who"].get("email") or split_emails(rec["邮箱"]) or [""])[0]
        else:
            kid, name = "", "（未匹配）" + (op["who"].get("kol") or "; ".join(op["who"].get("chat") or []))
            mail = (op["who"].get("email") or [""])[0]
        i = self.t.append()
        self.t.write_row(i, {"时间": to_serial(op["time"]), "ID": text_cell(kid), "KOL": text_cell(name),
                             "方向": op["dir"], "类型": text_cell(op["type"]), "摘要": text_cell(op["summary"]),
                             "渠道": text_cell(op["channel"]), "邮箱": text_cell(mail),
                             "来源": text_cell(source or "未注明"), "键": text_cell(key)})
        out = {"status": "added", "key": key, "matched": bool(i_kol), "id": kid}
        warning = "" if i_kol else master.chat_warning(op["who"], "按聊天账号找不到人")
        if warning:
            out["warning"] = warning
        return out


def is_auto_reply(etype) -> bool:
    """自动回复, KOC自动回复 and the like: machine messages that say nothing about whose turn it is."""
    return str(etype or "").strip().endswith("自动回复")


def ensure_project(book: Book, project: str):
    with contextlib.suppress(Exception):
        t = book.table(DICT_SHEET, PROJECT_TABLE)
        names = [str(r[0]).strip() for r in t.values() if r and r[0]]
        if project not in names:
            i = t.append()
            t.set(i, t.headers[0], text_cell(project))


def apply_event(op: dict, master: Master, log: Log, source: str) -> dict:
    """Rule-based bookkeeping for a single mail or chat message: move 待办方 / 下一步 / 阶段, then log it.

    A message already in the log only gets its summary refreshed (kol-mail sends the Chinese gist once the
    translation is ready). The state change comes before the log row, so a retry after an interruption can't skip it.

    A chat message (the op carries chat) is only ever booked on an existing row: kol-mail reports only chats linked to
    a row, so one that matches nothing points at a row deleted since, and must not leave an unmatched log line."""
    chat = bool(op["who"].get("chat"))
    if chat and not master.find(op["who"]):
        out = {"status": "not_found", "who": op["who"], "note": "总表里找不到这个聊天对应的 KOL，没有记"}
        warning = master.chat_warning(op["who"], "按聊天账号找不到人")
        if warning:
            out["warning"] = warning
        return out
    if log.find(log.key_for(op)) or is_auto_reply(op["type"]):
        return log.add({**op, "update_summary": True}, master, source)
    i = master.find(op["who"])
    if not i:
        return {**log.add(op, master, source), "note": "邮箱不在总表里，只记进了沟通记录"}
    rec_row = master.rows[i - 1]
    last_d = master.latest_contact(i)
    stage = master.t.get(rec_row, "阶段") or ""
    fields: dict = {}
    if last_d is None or op["time"] >= last_d - STALE_TOLERANCE:  # an older mail must not undo a later decision
        if op["dir"] == "收":
            fields["owner"] = "我方"
            fields["next"] = (f"{op['time']:%m-%d %H:%M} 收到 {op['channel']} 新消息，待处理" if chat
                              else f"{op['time']:%m-%d %H:%M} 收到新回信，待处理")
            if stage in ("未联系", "已发信", "退信", ""):
                fields["stage"] = "洽谈中"
        else:
            fields["owner"] = "对方"
            fields["next"] = (f"等对方回复（{op['time']:%m-%d} 已在 {op['channel']} 回复）" if chat
                              else f"等对方回复（{op['time']:%m-%d} 已回信）")
            if stage in ("未联系", ""):
                fields["stage"] = "已发信"
    update = master.update(i, fields, source)
    return {**log.add(op, master, source), "kol_update": update}


def apply_batch_row(op: dict, book: Book) -> dict:
    t = book.table(BATCH_SHEET, BATCH_TABLE)
    t.need(BATCH_HEADERS)
    if op["key"]:
        for row in t.values():
            if str(t.get(row, "Message-ID") or "") == op["key"]:
                return {"status": "exists"}
    i = t.append()
    t.write_row(i, {"批次": text_cell(op["batch"]), "项目": text_cell(op["project"]), "语言": text_cell(op["lang"]),
                    "日期": to_serial(op["date"]), "人数": op["count"], "主题": text_cell(op["subject"]),
                    "Message-ID": text_cell(op["key"])})
    return {"status": "added"}


def apply_ops(book: Book, ops: list[dict], source: str) -> list[dict]:
    master = Master(book)
    log = Log(book)
    results = []
    for op in ops:
        src = op.get("source") or source
        try:
            if op["op"] == "upsert":
                results.append(master.upsert(op["fields"], src, op["create"], op.get("rename", False),
                                             op.get("only_if_stage")))
            elif op["op"] == "log":
                results.append(log.add(op, master, src))
            elif op["op"] == "event":
                results.append(apply_event(op, master, log, src))
            elif op["op"] == "add_batch":
                results.append(apply_batch_row(op, book))
        except UsageError as e:  # e.g. an ambiguous name: report it and carry on with the others
            results.append({"status": "error", "op": op["op"], "error": str(e)})
    return results


# ---------------------------------------------------------------- running with lock, retry and queue

@contextlib.contextmanager
def file_lock(path: Path, wait: float, what: str):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    f = open(path, "a+")
    deadline = time.monotonic() + wait
    while True:
        try:
            f.seek(0)
            if sys.platform == "win32":
                msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                fcntl.flock(f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            break
        except OSError:
            if time.monotonic() >= deadline:
                f.close()
                raise Unavailable(what)
            time.sleep(0.3)
    try:
        yield
    finally:
        with contextlib.suppress(OSError):
            f.seek(0)
            if sys.platform == "win32":
                msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(f.fileno(), fcntl.LOCK_UN)
        f.close()


def process_lock():
    return file_lock(LOCKFILE, LOCK_WAIT_SECONDS, "另一个 kol.py 进程一直没结束")


def pending_lock():
    return file_lock(STATE_DIR / ".pending.lock", 30, "排队文件一直被占用")


def read_pending() -> tuple[list[dict], int]:
    """Queued entries and the number of lines read, so only those get removed after they are written."""
    with pending_lock():
        if not PENDING.exists():
            return [], 0
        lines = PENDING.read_text(encoding="utf-8").split("\n")
    entries = []
    for line in lines:
        if line.strip():
            try:
                entries.append(json.loads(line))
            except json.JSONDecodeError:
                with FAILED.open("a", encoding="utf-8") as f:
                    f.write(json.dumps({"error": "排队文件里有一行读不出来", "raw": line}, ensure_ascii=False) + "\n")
    return entries, len(lines)


def trim_pending(done_lines: int):
    """Drop the first done_lines lines; anything queued meanwhile by another process stays."""
    if not done_lines:
        return
    with pending_lock():
        if not PENDING.exists():
            return
        rest = PENDING.read_text(encoding="utf-8").split("\n")[done_lines:]
        if any(line.strip() for line in rest):
            PENDING.write_text("\n".join(rest) + "\n", encoding="utf-8")
        else:
            PENDING.unlink()


def queue(ops: list[dict], source: str, reason: str):
    with pending_lock():
        with PENDING.open("a", encoding="utf-8") as f:
            f.write(json.dumps({"queued_at": now().isoformat(), "source": source, "reason": reason,
                                "ops": [to_jsonable(o) for o in ops]}, ensure_ascii=False) + "\n")


def record_failed(ops: list[dict], source: str, error: str):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    with FAILED.open("a", encoding="utf-8") as f:
        f.write(json.dumps({"failed_at": now().isoformat(), "source": source, "error": error,
                            "ops": [to_jsonable(o) for o in ops]}, ensure_ascii=False) + "\n")


def run(ops: list[dict], source: str) -> dict:
    """Apply ops, or queue them; an unexpected error keeps them in pending.failed.jsonl so nothing is lost."""
    try:
        return _run(ops, source)
    except Exception as e:
        if ops:
            record_failed(ops, source, f"{type(e).__name__}: {e}")
        raise


def _run(ops: list[dict], source: str) -> dict:
    """Apply ops (already normalised) in one Excel session, earlier queued ops first. Queue everything if Excel can't."""
    started = time.monotonic()
    deadline = started + BUSY_WAIT_SECONDS
    reason = ""
    try:
        with process_lock():
            while True:
                try:
                    with open_book() as book:
                        book.backup_once_a_day()
                        flushed, failed = 0, []
                        entries, done_lines = read_pending()
                        for entry in entries:
                            entry_ops = [from_jsonable(o) for o in entry.get("ops", [])]
                            try:
                                done = apply_ops(book, entry_ops, entry.get("source", ""))
                                flushed += len(entry_ops)
                                bad = [(o, r) for o, r in zip(entry_ops, done)
                                       if r.get("status") in ("error", "needs_review")]
                                if bad:
                                    failed.append({**entry, "ops": [to_jsonable(o) for o, _ in bad],
                                                   "error": "; ".join(r["error"] for _, r in bad)})
                            except UsageError as e:  # e.g. a column was renamed since; keep it for a human
                                failed.append({**entry, "error": str(e)})
                        results = apply_ops(book, ops, source)
                        book.save()
                        mode = book.mode
                    if failed:
                        with FAILED.open("a", encoding="utf-8") as f:
                            for entry in failed:
                                f.write(json.dumps(entry, ensure_ascii=False) + "\n")
                    trim_pending(done_lines)
                    return {"status": "ok", "mode": mode, "results": results, "flushed_from_queue": flushed,
                            "queue_failed": len(failed)}
                except Unavailable as e:
                    reason = str(e)  # read-only or locked: won't clear up in seconds, so don't wait the full minute
                    deadline = min(deadline, started + UNAVAILABLE_WAIT_SECONDS)
                except Exception as e:  # noqa: BLE001
                    if is_com_error(e, BUSY_HRESULTS):
                        reason = "Excel 正忙（可能正在编辑单元格或开着对话框）"
                    elif is_com_error(e, GONE_HRESULTS):
                        reason = "Excel 在写入途中被关掉了"
                    else:
                        raise
                # >=: before 3.13 Windows' clock ticks every 15.6 ms, so with no wait at all ">" would retry anyway
                if time.monotonic() >= deadline:
                    break
                time.sleep(2)
    except Unavailable as e:
        reason = str(e)
    if ops:
        queue(ops, source, reason)
    return {"status": "queued", "reason": reason, "queued": len(ops),
            "note": "改动已排队，下次任何一次 kol.py 调用都会先补写；也可以手动运行 python kol.py flush"}


# ---------------------------------------------------------------- reading

def snapshot() -> tuple[list[dict], list[dict]]:
    """Current rows of the KOL and LOG tables. Live from Excel if the workbook is open there, else from the file.
    The Google sheet is always read live."""
    if GSHEET_ID or find_open_workbook_safe(WORKBOOK):
        with open_book(write=False) as book:
            m = Master(book)
            kols = [m.record(i) for i in range(1, len(m.rows) + 1)]
            lt = book.table(LOG_SHEET, LOG_TABLE)
            logs = []
            for row in lt.values():
                d = {h: lt.get(row, h) for h in lt.headers}
                d["时间"] = fmt_when(from_serial(d["时间"])) if isinstance(d["时间"], (int, float)) else (d["时间"] or "")
                logs.append({k: ("" if v is None else v) for k, v in d.items()})
            return kols, logs
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # "Conditional Formatting extension is not supported" (data bars)
        wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    def rows_of(sheet, table):
        ws = wb[sheet]
        ref = ws.tables[table].ref
        cells = list(ws[ref])
        headers = [c.value for c in cells[0]]
        out = []
        for line in cells[1:]:
            d = {}
            for h, c in zip(headers, line):
                v = c.value
                if isinstance(v, dt.datetime):
                    v = fmt_when(v)
                elif h in ("时间", "首次发信", "更新时间") and isinstance(v, (int, float)):
                    v = fmt_when(from_serial(v))
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                d[h] = "" if v is None else v
            if any(v not in ("", None) for k, v in d.items() if k not in FORMULA_COLUMNS):
                out.append(d)
        return out
    return rows_of(MAIN_SHEET, MAIN_TABLE), rows_of(LOG_SHEET, LOG_TABLE)


def find_open_workbook_safe(path: Path) -> bool:
    import pythoncom
    pythoncom.CoInitialize()
    try:
        wb = find_open_workbook(path)
        return wb is not None and bool(wb.Application.Visible or wb.Application.UserControl)
    except Exception:
        return False
    finally:
        pythoncom.CoUninitialize()


def match_record(rec: dict, who: dict) -> bool:
    if who.get("id") and str(rec.get("ID", "")).upper() == who["id"].upper():
        return True
    if who.get("email") and set(who["email"]) & set(split_emails(rec.get("邮箱"))):
        return True
    if who.get("chat") and {parse_chat(t)[0] for t in who["chat"]} & set(chat_ids(rec.get(CHAT_HEADER))):
        return True
    if who.get("kol") and who["kol"].casefold() in str(rec.get("KOL", "")).casefold():
        return True
    return False


# ---------------------------------------------------------------- CLI

def emit(obj, code: int = 0):
    print(json.dumps(obj, ensure_ascii=False, indent=2))
    sys.exit(code)


def add_kol_fields(p: argparse.ArgumentParser):
    for field, header in FIELDS.items():
        p.add_argument(f"--{field}", help=header)


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(prog="kol.py", description="KOL 总表写入工具（总表开着也能写）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("upsert", help="新建或更新一位 KOL（按 id / email / chat / channel / kol 名定位）")
    add_kol_fields(p)
    p.add_argument("--last", help="双方最近一封邮件的时间，只用来防止旧判断覆盖新情况，不写进表")
    p.add_argument("--no-create", action="store_true", help="找不到就不新建")
    p.add_argument("--source", default="手动会话", help="更新来源，写进“更新来源”列")

    for name, helptext in (("log", "往沟通记录里记一笔，不改 KOL 状态"),
                           ("event", "记一封收/发的邮件或一条聊天消息，并按规则更新待办方、下一步、阶段")):
        p = sub.add_parser(name, help=helptext)
        for f in ("id", "email", "kol", "chat", "dir", "type", "summary", "subject", "time", "message-id"):
            p.add_argument(f"--{f}")
        p.add_argument("--channel", help="渠道：邮件（默认）、Telegram、Discord")
        p.add_argument("--source", default="手动会话")

    p = sub.add_parser("batch", help="一次执行多条操作：JSON 文件路径，或 - 表示从标准输入读")
    p.add_argument("file")
    p.add_argument("--source", default="")
    p.add_argument("--skip-invalid", action="store_true", help="跳过格式不对的操作，其余照常写入")

    p = sub.add_parser("add-batch", help="记一次群发批次")
    for f in ("batch", "project", "lang", "date", "count", "subject", "message-id"):
        p.add_argument(f"--{f}")

    p = sub.add_parser("get", help="查一位 KOL 的整行和最近的沟通记录")
    for f in ("id", "email", "kol", "channel", "chat"):
        p.add_argument(f"--{f}")
    p.add_argument("--logs", type=int, default=10)

    p = sub.add_parser("list", help="按条件列出 KOL")
    for f in ("stage", "owner", "intent", "project", "lang"):
        p.add_argument(f"--{f}")
    p.add_argument("--fields", default="ID,KOL,阶段,意向,待办方,下一步")

    p = sub.add_parser("stats", help="各阶段、意向、待办人数")
    p.add_argument("--project")
    sub.add_parser("flush", help="把排队的改动写进总表")
    sub.add_parser("status", help="总表位置、是否打开、排队数")

    a = ap.parse_args(argv)
    try:
        if a.cmd == "upsert":
            raw = {f: getattr(a, f) for f in list(FIELDS) + ["last"] if getattr(a, f) is not None}
            ops = [norm_op({"op": "upsert", **raw, "create": not a.no_create})]
            res = run(ops, a.source)
        elif a.cmd in ("log", "event"):
            raw = {k: getattr(a, k.replace("-", "_")) for k in
                   ("id", "email", "kol", "chat", "dir", "type", "summary", "subject", "time", "channel")}
            raw["message_id"] = a.message_id
            ops = [norm_op({"op": a.cmd, **{k: v for k, v in raw.items() if v is not None}})]  # None = left out
            res = run(ops, a.source)
        elif a.cmd == "batch":
            raw_bytes = sys.stdin.buffer.read() if a.file == "-" else Path(a.file).read_bytes()
            data = json.loads(raw_bytes.decode("utf-8-sig"))
            items = data.get("ops", []) if isinstance(data, dict) else data
            source = a.source or (data.get("source", "") if isinstance(data, dict) else "") or "批量"
            ops, rejected = [], []
            for item in items:
                try:
                    ops.append(norm_op(item))
                except UsageError as e:
                    if not a.skip_invalid:
                        raise
                    rejected.append({"op": item, "error": str(e)})
            res = run(ops, source) if ops else {"status": "ok", "results": []}
            if rejected:
                res["rejected"] = rejected
        elif a.cmd == "add-batch":
            ops = [norm_op({"op": "add_batch", "batch": a.batch, "project": a.project, "lang": a.lang, "date": a.date,
                            "count": a.count, "subject": a.subject, "message_id": a.message_id})]
            res = run(ops, "群发记录")
        elif a.cmd == "flush":
            pending = sum(len(e.get("ops", [])) for e in read_pending()[0])
            if not pending:
                emit({"status": "ok", "note": "没有排队的改动"})
            res = run([], "补写队列")
        elif a.cmd == "status" and GSHEET_ID:
            emit({"workbook": f"https://docs.google.com/spreadsheets/d/{GSHEET_ID}/edit", "google_sheet": True,
                  "queued_ops": sum(len(e.get("ops", [])) for e in read_pending()[0]),
                  "failed_queue_file": str(FAILED) if FAILED.exists() else ""})
        elif a.cmd == "status":
            emit({"workbook": str(WORKBOOK), "exists": WORKBOOK.exists(),
                  "open_in_excel": find_open_workbook_safe(WORKBOOK) if WORKBOOK.exists() else False,
                  "queued_ops": sum(len(e.get("ops", [])) for e in read_pending()[0]),
                  "failed_queue_file": str(FAILED) if FAILED.exists() else ""})
        elif a.cmd == "get":
            who = norm_fields({k: getattr(a, k) for k in ("id", "email", "kol", "channel", "chat") if getattr(a, k)})
            if not who:
                raise UsageError("get 需要 --id / --email / --chat / --channel / --kol 其中一个")
            kols, logs = snapshot()
            hits = [r for r in kols if match_record(r, who)
                    or (who.get("channel") and str(r.get("频道链接", "")).rstrip("/").lower() == who["channel"].rstrip("/").lower())]
            ids = {h["ID"] for h in hits}
            mine = sorted((l for l in logs if l.get("ID") in ids), key=lambda l: str(l.get("时间", "")), reverse=True)
            emit({"matches": hits, "recent_logs": mine[:a.logs]})
        elif a.cmd == "list":
            kols, _ = snapshot()
            cond = {"阶段": a.stage, "待办方": a.owner, "意向": a.intent, "项目": a.project, "语言": a.lang}
            fields = [f.strip() for f in a.fields.split(",") if f.strip()]
            rows = [{f: r.get(f, "") for f in fields} for r in kols
                    if all(v is None or str(r.get(k, "")) == v for k, v in cond.items())]
            emit({"count": len(rows), "rows": rows})
        elif a.cmd == "stats":
            kols, _ = snapshot()
            if a.project:
                kols = [r for r in kols if r.get("项目") == a.project]
            count = lambda key: {v: sum(1 for r in kols if r.get(key) == v) for v in sorted({r.get(key, "") for r in kols})}
            emit({"total": len(kols), "阶段": count("阶段"), "意向": count("意向"), "待办方": count("待办方"),
                  "项目": count("项目")})
        if res.get("status") != "ok":
            emit(res, 3)
        problems = res.get("rejected") or [r for r in res.get("results", [])
                                           if r.get("status") in ("error", "needs_review")]
        emit(res, 4 if problems else 0)
    except UsageError as e:
        emit({"status": "error", "error": str(e)}, 2)
    except Unavailable as e:
        if not GSHEET_ID:
            raise
        emit({"status": "unavailable", "reason": str(e)}, 3)  # a read (list/get/stats) while Google is unreachable


if __name__ == "__main__":
    main()
