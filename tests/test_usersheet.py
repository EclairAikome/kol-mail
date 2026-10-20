"""Colleagues' own KOL sheets: guessing the column mapping, and syncing mail events into whatever layout they use
without changing it (sheet/usersheet.py). Excel itself is replaced by openpyxl on a temporary file."""

import contextlib
import datetime as dt
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "sheet"))

import usersheet as us  # noqa: E402

T0 = dt.datetime(2026, 10, 20, 9, 0)


def english_sheet(path: Path) -> Path:
    """A colleague's layout: a title row, English headers, their own words for whose turn it is."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Creators"
    wb.create_sheet("Notes")
    ws.append(["Q4 outreach"])
    ws.append(["Name", "Platform", "Channel Link", "Email", "Last Contact", "Ball", "Progress", "Fee"])
    ws.append(["Ana", "YouTube", "https://www.youtube.com/@ana", "ana@gmail.com", dt.datetime(2026, 10, 1, 8, 0), "Them",
               "我方回信：sent brief", 300])
    ws.append(["Bo", "Twitch", "", "bo@x.com; bo2@x.com", None, "Me", "", None])
    ws["E3"].number_format = "dd/mm/yyyy"
    for c in ws[4]:
        c.fill = PatternFill("solid", fgColor="FFF2CC")
        c.font = Font(italic=True)
    wb.save(path)
    return path


MAP = {"sheet": "Creators", "header_row": 2,
       "columns": {"name": "Name", "channel": "Channel Link", "email": "Email", "last_contact": "Last Contact",
                   "owner": "Ball", "summary": "Progress"}}


@contextlib.contextmanager
def grid_on(path):
    g = us.XlsxGrid(path, write=True)
    yield g


def sync(path, ops, mapping=MAP, monkeypatch=None):
    return us.run([us.norm_op(o) for o in ops], "test", mapping, grid_factory=lambda: grid_on(path))


@pytest.fixture(autouse=True)
def state_dir(tmp_path, monkeypatch):
    for name, value in (("STATE_DIR", tmp_path / "state"), ("PENDING", tmp_path / "state" / "pending.jsonl"),
                        ("STATE_FILE", tmp_path / "state" / "usersheet_state.json")):
        monkeypatch.setattr(us, name, value)


def cells(path, row):
    ws = load_workbook(path)["Creators"]
    return [c.value for c in ws[row]]


def test_mapping_is_guessed_from_headers_and_content(tmp_path):
    info = us.inspect(english_sheet(tmp_path / "a.xlsx"), None, None)
    assert info["sheet"] == "Creators" and info["sheets"] == ["Creators", "Notes"]
    assert info["header_row"] == 2
    assert info["guess"] == MAP["columns"]


def test_chinese_headers_and_an_email_column_with_an_odd_name(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["达人", "主页", "联系方式", "跟进情况", "日期", "负责人"])
    ws.append(["小王", "https://www.twitch.tv/w", "w@example.com", "谈价中", dt.datetime(2026, 9, 1), "吉利"])
    wb.save(tmp_path / "b.xlsx")
    guess = us.inspect(tmp_path / "b.xlsx", None, None)["guess"]
    assert guess["email"] == "联系方式"  # found by its content
    assert guess["name"] == "达人" and guess["channel"] == "主页" and guess["summary"] == "跟进情况"
    assert "owner" not in guess  # 负责人 is a person, not whose turn it is


def test_incoming_mail_updates_the_row_in_the_sheets_own_words(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    res = sync(path, [{"op": "event", "dir": "in", "email": "Ana@Gmail.com", "subject": "Re: brief",
                       "time": "2026-10-20T09:00:00", "message_id": "<m1>"}])
    assert res["status"] == "ok" and res["results"][0]["row"] == 3
    row = cells(path, 3)
    assert row[4] == T0 and row[5] == "Me" and row[6] == "对方来信：Re: brief"
    assert row[7] == 300  # columns that are not mapped are never touched
    assert load_workbook(path)["Creators"]["E3"].number_format == "dd/mm/yyyy"


def test_reply_and_second_address_in_one_cell(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    sync(path, [{"op": "event", "dir": "out", "email": "bo2@x.com", "summary": "我方回信：发了报价",
                 "time": "2026-10-20T09:00:00+08:00"}])
    row = cells(path, 4)
    assert row[5] == "Them" and row[6] == "我方回信：发了报价"


def test_unknown_kol_gets_a_new_row_styled_like_the_last(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    sync(path, [{"op": "event", "dir": "in", "email": "new@example.com", "name": "New KOL", "subject": "hi",
                 "time": "2026-10-20T09:00:00", "channel_url": "https://www.youtube.com/@newkol"}])
    ws = load_workbook(path)["Creators"]
    assert [c.value for c in ws[5]][:7] == ["New KOL", None, "https://www.youtube.com/@newkol", "new@example.com", T0,
                                            "Me", "对方来信：hi"]
    assert ws["A5"].font.italic and ws["A5"].fill.fgColor.rgb == ws["A4"].fill.fgColor.rgb


def test_new_row_extends_an_excel_table(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "KOL"
    ws.append(["KOL", "邮箱", "进展"])
    ws.append(["A", "a@x.com", ""])
    ws.add_table(Table(displayName="T", ref="A1:C2"))
    wb.save(tmp_path / "t.xlsx")
    m = {"sheet": "KOL", "header_row": 1, "columns": {"name": "KOL", "email": "邮箱", "summary": "进展"}}
    sync(tmp_path / "t.xlsx", [{"op": "event", "dir": "in", "email": "b@x.com", "subject": "s",
                                "time": "2026-10-20T09:00:00"}], m)
    ws = load_workbook(tmp_path / "t.xlsx")["KOL"]
    assert ws.tables["T"].ref == "A1:C3" and ws["B3"].value == "b@x.com" and ws["A3"].value == "b"


def test_an_older_mail_never_overwrites_newer_news(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    sync(path, [{"op": "event", "dir": "out", "email": "ana@gmail.com", "summary": "我方回信：新", "time": "2026-10-20T09:00:00"}])
    res = sync(path, [{"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "old", "time": "2026-10-19T09:00:00"}])
    assert "旧" in res["results"][0]["note"]
    assert cells(path, 3)[5:7] == ["Them", "我方回信：新"]


def test_same_mail_twice_only_refreshes_the_summary_with_the_chinese_gist(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    first = {"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "Re", "time": "2026-10-20T09:00:00",
             "message_id": "<m1>"}
    sync(path, [first])
    assert sync(path, [first])["results"][0]["status"] == "skipped"
    sync(path, [{**first, "summary": "对方回信：同意 300 美元"}])
    assert cells(path, 3)[6] == "对方回信：同意 300 美元"


def test_text_that_looks_like_a_formula_stays_text(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    sync(path, [{"op": "event", "dir": "in", "email": "ana@gmail.com", "summary": "=HYPERLINK(\"x\")",
                 "time": "2026-10-20T09:00:00"}])
    sync(path, [{"op": "event", "dir": "in", "email": "bo@x.com", "summary": "=HYPERLINK(\"x\")",
                 "time": "2026-10-20T09:00:00"}])
    ws = load_workbook(path)["Creators"]
    assert ws["G4"].value == "=HYPERLINK(\"x\")" and ws["G4"].data_type == "s"  # Bo's empty cell


def test_saved_channel_fills_only_an_empty_cell(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    sync(path, [{"op": "upsert", "email": "bo@x.com", "channel": "https://www.twitch.tv/bo"},
                {"op": "upsert", "email": "ana@gmail.com", "channel": "https://www.youtube.com/@other"}])
    assert cells(path, 4)[2] == "https://www.twitch.tv/bo" and cells(path, 3)[2] == "https://www.youtube.com/@ana"


def test_renamed_column_is_reported_not_guessed(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    m = {**MAP, "columns": {**MAP["columns"], "summary": "Notes"}}
    with pytest.raises(us.UsageError, match="进展摘要（Notes）"):
        sync(path, [{"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "s"}], m)


def test_busy_workbook_queues_and_flush_writes_later(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")

    @contextlib.contextmanager
    def busy():
        raise us.Unavailable("Excel 正忙")
        yield
    op = us.norm_op({"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "late", "time": "2026-10-20T09:00:00"})
    assert us.run([op], "test", MAP, grid_factory=busy)["status"] == "queued"
    assert cells(path, 3)[6] == "我方回信：sent brief"
    res = us.run([], "flush", MAP, grid_factory=lambda: grid_on(path))
    assert res["flushed_from_queue"] == 1 and cells(path, 3)[6] == "对方来信：late"
    assert not us.PENDING.exists()


def test_list_gives_app_the_addresses_and_links(tmp_path, monkeypatch):
    monkeypatch.setattr(us, "WORKBOOK", english_sheet(tmp_path / "a.xlsx"))
    rows = us.list_rows(MAP, ["邮箱", "频道链接", "阶段"])
    assert rows == [{"邮箱": "ana@gmail.com", "频道链接": "https://www.youtube.com/@ana", "阶段": ""},
                    {"邮箱": "bo@x.com; bo2@x.com", "频道链接": "", "阶段": ""}]


def test_new_sheet_is_mapped_and_takes_rows(tmp_path):
    made = us.create(tmp_path / "新表.xlsx")
    info = us.inspect(tmp_path / "新表.xlsx", None, None)
    assert info["guess"] == made["map"]["columns"]
    sync(tmp_path / "新表.xlsx", [{"op": "event", "dir": "in", "email": "a@x.com", "subject": "s",
                                   "time": "2026-10-20T09:00:00"}], made["map"])
    ws = load_workbook(tmp_path / "新表.xlsx")["KOL总表"]
    assert ws["D2"].value == "a@x.com" and ws["E2"].value == T0 and ws.tables["KOL"].ref == "A1:H2"
    with pytest.raises(us.UsageError):
        us.create(tmp_path / "新表.xlsx")


def test_mapping_must_have_the_email_column():
    with pytest.raises(us.UsageError, match="邮箱"):
        us.load_map(json.dumps({"sheet": "S", "columns": {"name": "Name"}}))


def test_automatic_replies_leave_the_row_alone_and_mail_links_fill_empty_cells(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    res = sync(path, [{"op": "event", "dir": "out", "email": "bo@x.com", "summary": "非工作时间自动回复（英语）",
                       "type": "自动回复", "time": "2026-10-20T09:00:00"},
                      {"op": "event", "dir": "in", "email": "bo@x.com", "subject": "hi", "time": "2026-10-20T09:00:00",
                       "channel_url": "https://www.twitch.tv/bo"},
                      {"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "hi",
                       "time": "2026-10-20T09:00:00", "channel_url": "https://www.youtube.com/@someone-else"}])
    assert res["results"][0]["status"] == "skipped"
    assert cells(path, 4)[2] == "https://www.twitch.tv/bo" and cells(path, 4)[6] == "对方来信：hi"
    assert cells(path, 3)[2] == "https://www.youtube.com/@ana"


def test_whose_turn_words_are_read_before_the_batch_changes_them(tmp_path):
    """Ana's "Them" becomes "Me" first in the same batch; Bo must still get the sheet's own word for their turn."""
    path = english_sheet(tmp_path / "a.xlsx")
    sync(path, [{"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "a", "time": "2026-10-20T09:00:00"},
                {"op": "event", "dir": "out", "email": "bo@x.com", "summary": "我方回信：b", "time": "2026-10-20T09:00:00"}])
    assert cells(path, 3)[5] == "Me" and cells(path, 4)[5] == "Them"
    # later the column shows only "Me" and "Them" is gone from it; the remembered word is still used
    sync(path, [{"op": "event", "dir": "in", "email": "bo@x.com", "subject": "c", "time": "2026-10-20T10:00:00"},
                {"op": "event", "dir": "out", "email": "ana@gmail.com", "summary": "我方回信：d", "time": "2026-10-20T10:00:00"}])
    assert cells(path, 4)[5] == "Me" and cells(path, 3)[5] == "Them"


def test_what_a_colleague_wrote_is_never_overwritten(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    ws_book = load_workbook(path)
    ws = ws_book["Creators"]
    ws["G3"], ws["F3"], ws["E4"] = "waiting for her manager", "Lisa", "next week"  # their own notes in mapped columns
    ws_book.save(path)
    res = sync(path, [{"op": "event", "dir": "in", "email": "ana@gmail.com", "subject": "Re", "time": "2026-10-20T09:00:00"},
                      {"op": "event", "dir": "out", "email": "bo@x.com", "summary": "我方回信：x", "time": "2026-10-20T09:00:00"}])
    assert "没有覆盖" in res["results"][0]["note"] and "最近联系时间" in res["results"][1]["note"]
    row3, row4 = cells(path, 3), cells(path, 4)
    assert row3[4] == T0 and row3[5] == "Lisa" and row3[6] == "waiting for her manager"
    assert row4[4] == "next week" and row4[5] == "对方" and row4[6] == "我方回信：x"  # no "Them" left in the sheet


def test_a_new_kol_never_lands_on_a_row_used_by_other_columns(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    book = load_workbook(path)
    book["Creators"]["H5"] = "=SUM(H3:H4)"  # a total right under the data, only in an unmapped column
    book.save(path)
    sync(path, [{"op": "event", "dir": "in", "email": "new@example.com", "subject": "hi", "time": "2026-10-20T09:00:00"}])
    ws = load_workbook(path)["Creators"]
    assert ws["H5"].value == "=SUM(H3:H4)" and ws["D5"].value is None
    assert ws["D6"].value == "new@example.com" and ws["A6"].font.italic


def test_a_text_formatted_date_column_gets_text(tmp_path):
    path = english_sheet(tmp_path / "a.xlsx")
    book = load_workbook(path)
    book["Creators"]["E4"].number_format = "@"
    book.save(path)
    sync(path, [{"op": "event", "dir": "in", "email": "bo@x.com", "subject": "hi", "time": "2026-10-20T09:00:00"}])
    assert cells(path, 4)[4] == "2026-10-20 09:00"


def test_people_and_status_columns_are_not_guessed_as_turn_or_progress(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Email", "Channel", "Status", "Notes", "Owner", "Date"])
    ws.append(["Ana", "ana@gmail.com", "https://www.youtube.com/@ana", "Negotiating", "likes CPS", "Lisa", None])
    wb.save(tmp_path / "c.xlsx")
    guess = us.inspect(tmp_path / "c.xlsx", None, None)["guess"]
    assert "owner" not in guess and "summary" not in guess and guess["email"] == "Email"
