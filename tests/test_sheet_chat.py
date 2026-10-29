"""Telegram / Discord accounts in the KOL sheets (聊天账号, tokens like tg:5550001234(@tea_cat)): kol.py's fixed layout
(here on the Google backend, against test_gsheet's fake Sheets API) and colleagues' own layouts (usersheet.py, on
openpyxl). A chat op from kol-mail carries id and chat, and when the linked row has them also email and kol."""

import contextlib
import importlib
import io
import json
import sys
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "sheet"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import gsheet  # noqa: E402
import usersheet as us  # noqa: E402
from test_gsheet import FakeSheets, kol_rows, log_rows  # noqa: E402  the fake Sheets API, not its tests

PIXELFOX, TIM = "dc:1300000000000000003", "tg:5550001234"
CHAT_EVENT = {"op": "event", "id": "K0001", "chat": f"{PIXELFOX}(PixelFox_)", "dir": "in", "channel": "Discord",
              "type": "Discord", "summary": "对方消息：可以合作", "time": "2026-10-29T14:12:27+08:00",
              "message_id": "dc:1100000000000000001:1400000000000000001"}


# ---------------------------------------------------------------- kol.py (fixed layout), Google backend

@pytest.fixture
def gsheet_kol(tmp_path, monkeypatch):
    fake = FakeSheets()
    monkeypatch.setattr(gsheet, "client_factory", lambda: httpx.Client(transport=httpx.MockTransport(fake.handler)))
    gsheet.Credentials(tmp_path).save({"client_id": "cid", "client_secret": "cs", "refresh_token": "r1"})
    api = gsheet.Api(tmp_path)
    sid = gsheet.create(api, "KOL总表（KOL Mail）")
    api.close()
    monkeypatch.setenv("KOL_SHEET_PATH", f"gsheet:{sid}")
    monkeypatch.setenv("KOL_SHEET_STATE", str(tmp_path))
    monkeypatch.syspath_prepend(str(ROOT / "sheet"))
    kol = importlib.reload(importlib.import_module("kol"))
    monkeypatch.setattr(kol, "UNAVAILABLE_WAIT_SECONDS", 0)
    return kol, fake, sid


def run(kol, *ops):
    res = kol.run([kol.norm_op(o) for o in ops], "test")
    assert res["status"] == "ok", res
    return res["results"]


def drop_chat_column(fake, sid):
    """A sheet made before 2026-10-29: 聊天账号 is the last header and nothing is under it yet."""
    header = fake.tab(sid, "KOL总表")["cells"][0]
    assert header[-1] == "聊天账号"
    header.pop()


def cli(kol, capsys, *argv):
    with pytest.raises(SystemExit) as stop:
        kol.main(list(argv))
    return stop.value.code, json.loads(capsys.readouterr().out)


def test_new_sheets_have_the_column_and_old_ones_still_pass_the_check(gsheet_kol, tmp_path):
    kol, fake, sid = gsheet_kol
    assert fake.rows(sid, "KOL总表")[0][-1] == "聊天账号" == kol.MAIN_HEADERS[-1]
    drop_chat_column(fake, sid)
    api = gsheet.Api(tmp_path)
    assert gsheet.check(api, sid) == []
    fake.tab(sid, "KOL总表")["cells"][0].remove("邮箱")
    assert gsheet.check(api, sid) == ["“KOL总表”页缺少列：邮箱"]
    api.close()


def test_chat_accounts_merge_and_find_the_row(gsheet_kol, capsys):
    kol, fake, sid = gsheet_kol
    run(kol, {"op": "upsert", "kol": "PixelFox", "email": "pixelfox@x.com", "stage": "已发信"})
    res = run(kol, {"op": "upsert", "id": "K0001", "chat": f"{PIXELFOX}(PixelFox_)", "create": False},
              {"op": "upsert", "id": "K0001", "chat": f"{TIM}(@tea_cat)", "create": False},
              {"op": "upsert", "id": "K0001", "chat": f"{PIXELFOX}(Pixel)", "create": False},
              {"op": "upsert", "id": "K0001", "chat": PIXELFOX, "create": False})
    assert [r["status"] for r in res] == ["updated", "updated", "updated", "unchanged"]
    assert kol_rows(fake, sid)[0]["聊天账号"] == f"{PIXELFOX}(Pixel); {TIM}(@tea_cat)"

    lone = {k: v for k, v in CHAT_EVENT.items() if k != "id"}  # found by the account alone
    res = run(kol, lone)
    assert res[0]["status"] == "added" and res[0]["id"] == "K0001"
    row = kol_rows(fake, sid)[0]
    assert row["待办方"] == "我方" and row["阶段"] == "洽谈中" and "收到 Discord 新消息" in row["下一步"]
    log = log_rows(fake, sid)[0]
    assert (log["渠道"], log["类型"], log["键"], log["邮箱"]) == ("Discord", "Discord", lone["message_id"], "pixelfox@x.com")
    assert run(kol, {**lone, "summary": "对方消息：可以合作，100 美元"})[0]["status"] == "summary_updated"
    assert len(log_rows(fake, sid)) == 1

    code, out = cli(kol, capsys, "get", "--chat", "TG:5550001234")
    assert code == 0 and [m["ID"] for m in out["matches"]] == ["K0001"]
    code, out = cli(kol, capsys, "list", "--fields", "ID,邮箱,聊天账号")
    assert out["rows"] == [{"ID": "K0001", "邮箱": "pixelfox@x.com", "聊天账号": f"{PIXELFOX}(Pixel); {TIM}(@tea_cat)"}]


def test_a_chat_op_with_the_rows_email_and_name_and_no_email_at_all(gsheet_kol):
    kol, fake, sid = gsheet_kol
    run(kol, {"op": "upsert", "kol": "PixelFox", "email": "pixelfox@x.com"}, {"op": "upsert", "kol": "Tim"})
    res = run(kol, {**CHAT_EVENT, "email": "pixelfox@x.com", "kol": "PixelFox"},
              {"op": "event", "id": "K0002", "chat": f"{TIM}(@tea_cat)", "kol": "Tim", "dir": "out",
               "channel": "Telegram", "type": "Telegram", "summary": "我方消息：发了 brief",
               "time": "2026-10-29T15:00:00+08:00", "message_id": "tg:5550001234:7"})
    assert [r["id"] for r in res] == ["K0001", "K0002"]
    tim = kol_rows(fake, sid)[1]
    assert tim["待办方"] == "对方" and "已在 Telegram 回复" in tim["下一步"] and tim["阶段"] == "已发信"
    assert log_rows(fake, sid)[1]["邮箱"] in (None, "")


def test_a_chat_that_matches_no_row_writes_nothing(gsheet_kol):
    kol, fake, sid = gsheet_kol
    run(kol, {"op": "upsert", "kol": "PixelFox", "email": "pixelfox@x.com"})
    res = run(kol, {**CHAT_EVENT, "id": "K0404", "kol": "PixelFox"})
    assert res[0]["status"] == "not_found"
    assert log_rows(fake, sid) == [] and kol_rows(fake, sid)[0]["待办方"] in (None, "")


def test_a_sheet_without_the_column_keeps_working(gsheet_kol):
    kol, fake, sid = gsheet_kol
    drop_chat_column(fake, sid)
    res = run(kol, {"op": "upsert", "kol": "PixelFox", "email": "pixelfox@x.com", "chat": PIXELFOX, "next": "等报价"},
              CHAT_EVENT, {**CHAT_EVENT, "id": None, "message_id": "dc:1:2"},
              {"op": "event", "dir": "in", "email": "pixelfox@x.com", "subject": "Re", "time": "2026-10-29T16:00:00",
               "message_id": "<m1>"})
    assert res[0]["status"] == "created" and "在“KOL总表”页表格最右边" in res[0]["warning"]
    assert res[1]["status"] == "added" and "warning" not in res[1]  # found by id
    assert res[2]["status"] == "not_found" and "聊天账号" in res[2]["warning"]
    assert res[3]["status"] == "added"
    assert [r["渠道"] for r in log_rows(fake, sid)] == ["Discord", "邮件"]
    assert "聊天账号" not in fake.rows(sid, "KOL总表")[0]


def test_invalid_tokens_are_refused_per_op(gsheet_kol, capsys, monkeypatch):
    kol, fake, sid = gsheet_kol
    run(kol, {"op": "upsert", "kol": "PixelFox", "email": "pixelfox@x.com"})
    batch = {"ops": [{"op": "upsert", "id": "K0001", "chat": "@tea_cat"}, CHAT_EVENT]}
    monkeypatch.setattr(sys, "stdin", type("S", (), {"buffer": io.BytesIO(json.dumps(batch).encode("utf-8"))})())
    code, out = cli(kol, capsys, "batch", "-", "--skip-invalid")
    assert code == 4 and "聊天账号格式不对" in out["rejected"][0]["error"]
    assert out["results"][0]["status"] == "added"


# ---------------------------------------------------------------- usersheet.py (a colleague's own layout)

@pytest.fixture(autouse=True)
def state_dir(tmp_path, monkeypatch):
    for name, value in (("STATE_DIR", tmp_path / "state"), ("PENDING", tmp_path / "state" / "pending.jsonl"),
                        ("STATE_FILE", tmp_path / "state" / "usersheet_state.json")):
        monkeypatch.setattr(us, name, value)


def chat_sheet(path: Path, chat_cells=("tg:5550001234(@tea_cat)", "", "my own note: ask on TG")) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Creators"
    ws.append(["Name", "Email", "TG", "Last Contact", "Ball", "Progress"])
    ws.append(["Tim", "", chat_cells[0], None, "Them", "我方消息：发了 brief"])
    ws.append(["Ana", "ana@gmail.com", chat_cells[1], None, "Them", ""])
    ws.append(["Bo", "bo@x.com", chat_cells[2], None, "", ""])
    wb.save(path)
    return path


MAP = {"sheet": "Creators", "header_row": 1,
       "columns": {"name": "Name", "email": "Email", "chat": "TG", "last_contact": "Last Contact", "owner": "Ball",
                   "summary": "Progress"}}
NO_CHAT = {**MAP, "columns": {k: v for k, v in MAP["columns"].items() if k != "chat"}}


@contextlib.contextmanager
def grid_on(path):
    yield us.XlsxGrid(path, write=True)


def sync(path, ops, mapping=MAP):
    return us.run([us.norm_op(o) for o in ops], "test", mapping, grid_factory=lambda: grid_on(path))["results"]


def cells(path, row):
    return [c.value for c in load_workbook(path)["Creators"][row]]


def test_chat_columns_are_guessed_without_taking_the_email_or_a_link_column(tmp_path):
    assert us.inspect(chat_sheet(tmp_path / "a.xlsx"), None, None)["guess"] == MAP["columns"]
    wb = Workbook()
    ws = wb.active
    ws.append(["KOL", "Email / Discord", "Telegram", "Last Game", "Discord"])
    ws.append(["A", "a@x.com", "https://t.me/a", "Elden Ring", "dc:1(a)"])
    wb.save(tmp_path / "b.xlsx")
    guess = us.inspect(tmp_path / "b.xlsx", None, None)["guess"]
    assert guess == {"name": "KOL", "email": "Email / Discord", "chat": "Discord"}
    assert us.FIELDS["chat"] == "聊天账号"


def test_chat_message_with_the_rows_email_fills_the_account_and_the_turn(tmp_path):
    path = chat_sheet(tmp_path / "a.xlsx")
    res = sync(path, [{**CHAT_EVENT, "email": "ana@gmail.com", "kol": "Ana", "time": "2026-10-29T09:00:00"}])
    assert res[0]["status"] == "ok" and res[0]["row"] == 3
    assert cells(path, 3)[2:] == [f"{PIXELFOX}(PixelFox_)", us.parse_when("2026-10-29T09:00:00"), "我方", "对方消息：可以合作"]


def test_chat_message_without_email_is_found_by_its_account_and_id_is_ignored(tmp_path):
    path = chat_sheet(tmp_path / "a.xlsx")
    res = sync(path, [{"op": "event", "id": "K0103", "chat": "TG:5550001234(@yummy_new)", "dir": "in",
                       "channel": "Telegram", "summary": "对方消息：报价 100", "time": "2026-10-29T09:00:00",
                       "message_id": "tg:5550001234:5"}])
    assert res[0]["row"] == 2
    row = cells(path, 2)
    assert row[2] == f"{TIM}(@yummy_new)"  # only accounts in the cell: the newer display name is taken
    assert row[4:] == ["我方", "对方消息：报价 100"]  # 我方消息… was written by this tool: replaced


def test_unknown_chat_gets_a_new_row_and_a_colleagues_text_is_kept(tmp_path):
    path = chat_sheet(tmp_path / "a.xlsx")
    res = sync(path, [{"op": "event", "chat": "dc:77(NewGuy)", "kol": "New Guy", "dir": "in", "summary": "对方消息：hi",
                       "time": "2026-10-29T09:00:00", "message_id": "dc:77:1"},
                      {"op": "event", "chat": "dc:88", "dir": "in", "summary": "对方消息：yo",
                       "time": "2026-10-29T09:00:00", "message_id": "dc:88:1"},
                      {"op": "event", "chat": "tg:5", "email": "bo@x.com", "dir": "out", "summary": "我方消息：ok",
                       "time": "2026-10-29T09:00:00", "message_id": "tg:5:1"}])
    assert [r.get("created") for r in res] == [True, True, False]
    assert cells(path, 5)[:3] == ["New Guy", None, "dc:77(NewGuy)"]
    assert cells(path, 6)[:3] == ["dc:88", None, "dc:88"]
    assert cells(path, 4)[2] == "my own note: ask on TG" and "聊天账号" in res[2]["note"]
    again = sync(path, [{"op": "event", "chat": "dc:77", "dir": "in", "summary": "对方消息：again",
                         "time": "2026-10-29T10:00:00", "message_id": "dc:77:2"}])
    assert again[0]["row"] == 5 and len(load_workbook(path)["Creators"]["A"]) == 6


def test_without_a_mapped_chat_column_a_chat_without_email_is_skipped_not_failed(tmp_path, monkeypatch, capsys):
    path = chat_sheet(tmp_path / "a.xlsx")
    ops = [{**CHAT_EVENT, "kol": "Tim"}, {"op": "upsert", "id": "K0001", "chat": TIM},
           {**CHAT_EVENT, "email": "ana@gmail.com", "message_id": "dc:1:9"}]
    monkeypatch.setenv("KOL_SHEET_MAP", json.dumps(NO_CHAT))
    real = us.run
    monkeypatch.setattr(us, "run", lambda o, s, m: real(o, s, m, grid_factory=lambda: grid_on(path)))
    (tmp_path / "ops.json").write_text(json.dumps({"ops": ops}), encoding="utf-8")
    with pytest.raises(SystemExit) as stop:
        us.main(["batch", str(tmp_path / "ops.json"), "--skip-invalid"])
    out = json.loads(capsys.readouterr().out)
    assert stop.value.code == 0
    assert [r["status"] for r in out["results"]] == ["skipped", "skipped", "ok"]
    assert cells(path, 3)[4:] == ["我方", "对方消息：可以合作"] and cells(path, 3)[2] in (None, "")  # TG not mapped
    assert len(load_workbook(path)["Creators"]["A"]) == 4


def test_upsert_records_a_linked_account_but_never_adds_a_row(tmp_path):
    path = chat_sheet(tmp_path / "a.xlsx")
    res = sync(path, [{"op": "upsert", "id": "K0001", "email": "ana@gmail.com", "chat": f"{PIXELFOX}(Pixel)",
                       "kol": "Ana"},
                      {"op": "upsert", "chat": f"{TIM}(@tea_cat)"},
                      {"op": "upsert", "chat": "dc:404"},
                      {"op": "upsert", "email": "bo@x.com", "chat": "tg:5"}])
    assert [r["status"] for r in res] == ["ok", "ok", "skipped", "ok"]
    assert res[0]["updated"] == ["chat"] and res[1]["updated"] == [] and "没有覆盖" in res[3]["note"]
    assert cells(path, 3)[2] == f"{PIXELFOX}(Pixel)" and cells(path, 4)[2] == "my own note: ask on TG"
    assert len(load_workbook(path)["Creators"]["A"]) == 4


def test_a_bad_token_is_refused_and_an_op_needs_email_or_chat():
    with pytest.raises(us.UsageError, match="聊天账号格式不对"):
        us.norm_op({**CHAT_EVENT, "chat": "@tea_cat"})
    with pytest.raises(us.UsageError, match="邮箱或聊天账号"):
        us.norm_op({"op": "event", "id": "K0001", "dir": "in", "summary": "x"})
    op = us.norm_op({"op": "event", "chat": TIM, "dir": "out", "time": "2026-10-29T09:00:00"})
    assert op["summary"] == "我方消息：（无内容）" and op["email"] == ""
    old = us.norm_op({"op": "event", "email": "a@x.com", "dir": "in", "subject": "hi"})
    assert old["summary"] == "对方来信：hi" and old["chat"] == []


def test_list_gives_app_the_chat_accounts(tmp_path, monkeypatch):
    monkeypatch.setattr(us, "RAW_WORKBOOK", str(chat_sheet(tmp_path / "a.xlsx")))
    rows = us.list_rows(MAP, ["ID", "邮箱", "聊天账号"])
    assert rows[0] == {"ID": "", "邮箱": "", "聊天账号": f"{TIM}(@tea_cat)"}
    assert us.list_rows(NO_CHAT, ["邮箱", "聊天账号"])[0] == {"邮箱": "ana@gmail.com", "聊天账号": ""}
